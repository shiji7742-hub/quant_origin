"""
回测：板块信号 + 涨停破位洗盘（多线程版本）
逻辑：当某月某板块出现潜力信号时，在该板块中找涨停破位洗盘的个股
"""
import akshare as ak
import pandas as pd
import numpy as np
from datetime import datetime, timedelta
from concurrent.futures import ThreadPoolExecutor, as_completed
from threading import Lock
import time
from strategies import strategy_limit_up_washout

# 2025年每月检测到的潜力板块信号（扩展版，每月5个板块）
MONTHLY_SECTOR_SIGNALS = {
    '2025-01': ['光刻机', '存储芯片', 'AIGC概念', '算力租赁', '英伟达概念'],
    '2025-02': ['DeepSeek概念', '机器人执行器', '减速器', '人形机器人', '谷歌概念'],
    '2025-03': ['低空经济', '无人机', '通用航空', '飞行汽车', 'eVTOL概念'],
    '2025-04': ['固态电池', '锂电池', '充电桩', '钠离子电池', '新能源汽车'],
    '2025-05': ['GDR', '电子后视镜', '乳业', '举牌', '3D摄像头'],
    '2025-06': ['人工智能', '大模型', 'ChatGPT概念', '数据要素', '云计算'],
    '2025-07': ['芯片', '半导体', '先进封装', 'Chiplet概念', '国产替代'],
    '2025-08': ['消费电子', '折叠屏', '华为概念', '苹果概念', 'MR/VR'],
    '2025-09': ['医药', '创新药', '医疗器械', 'CRO', '中药'],
    '2025-10': ['白酒', '食品饮料', '预制菜', '调味品', '乳制品'],
    '2025-11': ['银行', '保险', '证券', '金融科技', '数字货币'],
    '2025-12': ['光伏', '储能', '风电', '氢能源', '碳中和'],
}

# 线程安全的打印锁
print_lock = Lock()

def safe_print(msg):
    with print_lock:
        print(msg)

def get_sector_stocks(sector_name):
    """获取板块成分股"""
    try:
        df = ak.stock_board_concept_cons_em(symbol=sector_name)
        if df is None:
            return []
        df = df[df['代码'].str.match(r'^(60|00)')]
        df = df[~df['名称'].str.contains('ST')]
        return df[['代码', '名称']].head(15).to_dict('records')
    except:
        return []

def get_stock_data(symbol, start_date, end_date):
    """获取股票历史数据"""
    try:
        df = ak.stock_zh_a_hist(symbol=symbol, period="daily", 
                                start_date=start_date, end_date=end_date, adjust="qfq")
        if df is not None and len(df) > 0:
            df['日期'] = pd.to_datetime(df['日期'])
        return df
    except:
        return None

def find_washout_signals_in_month(df, month_start, month_end):
    """在指定月份内找涨停破位洗盘信号"""
    if df is None or len(df) < 30:
        return []
    
    signals = []
    df = df.reset_index(drop=True)
    
    for i in range(30, len(df)):
        current_date = df.iloc[i]['日期']
        if current_date < month_start or current_date > month_end:
            continue
        
        df_slice = df.iloc[:i+1].copy()
        result = strategy_limit_up_washout(df_slice, relaxed=True)
        
        if result['触发']:
            signals.append({
                'date': current_date,
                'price': df.iloc[i]['收盘'],
                'idx': i,
                'detail': result['说明']
            })
    
    return signals

def calculate_returns(df, signal_idx):
    """计算信号后的收益"""
    signal_price = df.iloc[signal_idx]['收盘']
    returns = {}
    
    for days in [5, 10, 20, 40, 60]:
        target_idx = signal_idx + days
        if target_idx < len(df):
            future_price = df.iloc[target_idx]['收盘']
            ret = (future_price - signal_price) / signal_price * 100
            returns[f'{days}日'] = ret
        else:
            returns[f'{days}日'] = None
    
    if signal_idx + 1 < len(df):
        future_data = df.iloc[signal_idx+1:min(signal_idx+41, len(df))]
        if len(future_data) > 0:
            max_price = future_data['最高'].max()
            min_price = future_data['最低'].min()
            returns['最大收益'] = (max_price - signal_price) / signal_price * 100
            returns['最大回撤'] = (min_price - signal_price) / signal_price * 100
    
    return returns

def process_stock(args):
    """处理单只股票（线程任务）"""
    code, name, sector, month_key, month_start, month_end, data_start, data_end = args
    
    try:
        df = get_stock_data(code, data_start, data_end)
        if df is None or len(df) < 50:
            return []
        
        signals = find_washout_signals_in_month(df, month_start, month_end)
        results = []
        
        for sig in signals:
            returns = calculate_returns(df, sig['idx'])
            ret_5 = returns.get('5日', 0) or 0
            ret_10 = returns.get('10日', 0) or 0
            
            safe_print(f"  ✓ {code} {name} [{sector}] {sig['date'].strftime('%m-%d')} "
                      f"5日:{ret_5:+.1f}% 10日:{ret_10:+.1f}%")
            
            results.append({
                '月份': month_key,
                '板块': sector,
                '代码': code,
                '名称': name,
                '信号日期': sig['date'].strftime('%Y-%m-%d'),
                '买入价': round(sig['price'], 2),
                '触发说明': sig['detail'],
                **returns
            })
        
        return results
    except Exception as e:
        return []

def backtest_sector_washout_parallel(max_workers=8):
    """多线程并行回测"""
    print("=" * 70)
    print("回测：板块信号 + 涨停破位洗盘（多线程版本）")
    print("=" * 70)
    print(f"\n使用 {max_workers} 个线程并行处理")
    
    start_time = time.time()
    
    # 收集所有任务
    all_tasks = []
    scanned_stocks = set()
    
    for month_key, sectors in MONTHLY_SECTOR_SIGNALS.items():
        year, month = month_key.split('-')
        month_start = pd.to_datetime(f"{year}-{month}-01")
        
        if int(month) == 12:
            month_end = pd.to_datetime(f"{int(year)+1}-01-01") - timedelta(days=1)
        else:
            month_end = pd.to_datetime(f"{year}-{int(month)+1:02d}-01") - timedelta(days=1)
        
        data_start = (month_start - timedelta(days=90)).strftime('%Y%m%d')
        data_end = (month_end + timedelta(days=70)).strftime('%Y%m%d')
        
        print(f"\n准备【{month_key}】: {', '.join(sectors)}")
        
        for sector in sectors:
            stocks = get_sector_stocks(sector)
            if not stocks:
                continue
            
            for stock in stocks:
                code = stock['代码']
                name = stock['名称']
                
                stock_month_key = f"{code}_{month_key}"
                if stock_month_key in scanned_stocks:
                    continue
                scanned_stocks.add(stock_month_key)
                
                all_tasks.append((code, name, sector, month_key, 
                                 month_start, month_end, data_start, data_end))
    
    print(f"\n总任务数: {len(all_tasks)} 只股票待检测")
    print("=" * 70)
    print("开始并行扫描...")
    
    # 并行执行
    all_signals = []
    completed = 0
    
    with ThreadPoolExecutor(max_workers=max_workers) as executor:
        futures = {executor.submit(process_stock, task): task for task in all_tasks}
        
        for future in as_completed(futures):
            completed += 1
            if completed % 50 == 0:
                safe_print(f"\n--- 进度: {completed}/{len(all_tasks)} ---\n")
            
            try:
                results = future.result()
                if results:
                    all_signals.extend(results)
            except Exception as e:
                pass
    
    elapsed = time.time() - start_time
    print(f"\n扫描完成！耗时: {elapsed:.1f}秒")
    
    return all_signals

def analyze_results(signals):
    """分析回测结果"""
    if not signals:
        print("\n没有找到任何信号")
        return None
    
    df = pd.DataFrame(signals)
    
    print("\n" + "=" * 70)
    print("回测结果分析")
    print("=" * 70)
    
    print(f"\n总信号数: {len(df)}")
    print(f"涉及月份: {df['月份'].nunique()}个")
    print(f"涉及板块: {df['板块'].nunique()}个")
    print(f"涉及个股: {df['代码'].nunique()}只")
    
    # 按持有期统计
    print("\n【不同持有期收益统计】")
    print("-" * 60)
    
    period_stats = []
    for period in ['5日', '10日', '20日', '40日', '60日']:
        if period in df.columns:
            valid = df[df[period].notna()][period]
            if len(valid) > 0:
                win_rate = (valid > 0).sum() / len(valid) * 100
                avg_ret = valid.mean()
                median_ret = valid.median()
                
                print(f"\n{period}:")
                print(f"  样本数: {len(valid)}, 胜率: {win_rate:.1f}%")
                print(f"  平均收益: {avg_ret:.2f}%, 中位数: {median_ret:.2f}%")
                print(f"  最大盈利: {valid.max():.2f}%, 最大亏损: {valid.min():.2f}%")
                
                period_stats.append({
                    '持有期': period,
                    '样本数': len(valid),
                    '胜率': f"{win_rate:.1f}%",
                    '平均收益': f"{avg_ret:.2f}%",
                    '中位数': f"{median_ret:.2f}%"
                })
    
    # 盈亏比
    if '最大收益' in df.columns and '最大回撤' in df.columns:
        valid_max = df[df['最大收益'].notna()]['最大收益']
        valid_dd = df[df['最大回撤'].notna()]['最大回撤']
        if len(valid_max) > 0 and valid_dd.mean() != 0:
            print(f"\n【40日内盈亏分析】")
            print(f"  平均最大收益: {valid_max.mean():.2f}%")
            print(f"  平均最大回撤: {valid_dd.mean():.2f}%")
            print(f"  盈亏比: {abs(valid_max.mean() / valid_dd.mean()):.2f}")
    
    # 按月份统计
    print("\n【按月份统计】")
    print("-" * 60)
    
    month_stats = []
    for month in sorted(df['月份'].unique()):
        month_df = df[df['月份'] == month]
        valid = month_df[month_df['10日'].notna()]['10日']
        if len(valid) > 0:
            win_rate = (valid > 0).sum() / len(valid) * 100
            avg_ret = valid.mean()
            print(f"  {month}: 信号{len(month_df)}个, 10日胜率{win_rate:.1f}%, 平均{avg_ret:+.2f}%")
            
            month_stats.append({
                '月份': month,
                '信号数': len(month_df),
                '10日胜率': f"{win_rate:.1f}%",
                '10日平均': f"{avg_ret:+.2f}%"
            })
    
    # 按板块统计
    print("\n【按板块统计】")
    print("-" * 60)
    
    sector_stats = []
    for sector in df['板块'].unique():
        sector_df = df[df['板块'] == sector]
        valid = sector_df[sector_df['10日'].notna()]['10日']
        if len(valid) > 0:
            win_rate = (valid > 0).sum() / len(valid) * 100
            avg_ret = valid.mean()
            print(f"  {sector}: 信号{len(sector_df)}个, 10日胜率{win_rate:.1f}%, 平均{avg_ret:+.2f}%")
            
            sector_stats.append({
                '板块': sector,
                '信号数': len(sector_df),
                '10日胜率': f"{win_rate:.1f}%",
                '10日平均': f"{avg_ret:+.2f}%"
            })
    
    return {
        'df': df,
        'period_stats': period_stats,
        'month_stats': month_stats,
        'sector_stats': sector_stats
    }

def print_all_cases(signals):
    """打印所有案例"""
    if not signals:
        return
    
    print("\n" + "=" * 70)
    print("所有案例明细")
    print("=" * 70)
    
    df = pd.DataFrame(signals)
    
    for month in sorted(df['月份'].unique()):
        month_df = df[df['月份'] == month].sort_values('信号日期')
        
        print(f"\n【{month}】共{len(month_df)}个信号")
        print("-" * 60)
        
        for _, row in month_df.iterrows():
            ret_5 = row.get('5日', 0) or 0
            ret_10 = row.get('10日', 0) or 0
            ret_20 = row.get('20日', 0) or 0
            max_ret = row.get('最大收益', 0) or 0
            max_dd = row.get('最大回撤', 0) or 0
            
            result = "✓盈" if ret_10 > 0 else "✗亏"
            
            print(f"\n  {row['代码']} {row['名称']} [{row['板块']}]")
            print(f"    信号日: {row['信号日期']}, 买入价: {row['买入价']}")
            print(f"    收益: 5日{ret_5:+.1f}% | 10日{ret_10:+.1f}% | 20日{ret_20:+.1f}% {result}")
            print(f"    40日内: 最高{max_ret:+.1f}%, 最低{max_dd:+.1f}%")

def export_to_excel(signals, results):
    """导出到Excel"""
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    
    filename = f"板块信号+涨停洗盘回测_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    wb = openpyxl.Workbook()
    
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    green_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )
    
    # Sheet1: 所有案例
    ws1 = wb.active
    ws1.title = "所有案例"
    
    headers = ['月份', '板块', '代码', '名称', '信号日期', '买入价', 
               '5日', '10日', '20日', '40日', '最大收益', '最大回撤', '10日结果']
    
    for col, h in enumerate(headers, 1):
        cell = ws1.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = border
    
    for row_idx, sig in enumerate(signals, 2):
        ret_10 = sig.get('10日', 0) or 0
        result = "盈利" if ret_10 > 0 else "亏损"
        
        row_data = [
            sig['月份'], sig['板块'], sig['代码'], sig['名称'],
            sig['信号日期'], sig['买入价'],
            f"{sig.get('5日', 0) or 0:.2f}%",
            f"{sig.get('10日', 0) or 0:.2f}%",
            f"{sig.get('20日', 0) or 0:.2f}%",
            f"{sig.get('40日', 0) or 0:.2f}%",
            f"{sig.get('最大收益', 0) or 0:.2f}%",
            f"{sig.get('最大回撤', 0) or 0:.2f}%",
            result
        ]
        
        for col, val in enumerate(row_data, 1):
            cell = ws1.cell(row=row_idx, column=col, value=val)
            cell.border = border
            
            if col in [7, 8, 9, 10, 11]:
                try:
                    num_val = float(str(val).replace('%', ''))
                    if num_val > 0:
                        cell.fill = green_fill
                    elif num_val < 0:
                        cell.fill = red_fill
                except:
                    pass
            
            if col == 13:
                cell.fill = green_fill if val == "盈利" else red_fill
    
    # Sheet2: 统计汇总
    ws2 = wb.create_sheet(title="统计汇总")
    
    row = 1
    ws2.cell(row=row, column=1, value="【持有期统计】").font = Font(bold=True, size=12)
    row += 2
    
    if results and 'period_stats' in results:
        for col, h in enumerate(['持有期', '样本数', '胜率', '平均收益', '中位数'], 1):
            cell = ws2.cell(row=row, column=col, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.border = border
        
        for stat in results['period_stats']:
            row += 1
            for col, h in enumerate(['持有期', '样本数', '胜率', '平均收益', '中位数'], 1):
                cell = ws2.cell(row=row, column=col, value=stat.get(h, ''))
                cell.border = border
    
    row += 3
    ws2.cell(row=row, column=1, value="【月份统计】").font = Font(bold=True, size=12)
    row += 2
    
    if results and 'month_stats' in results:
        for col, h in enumerate(['月份', '信号数', '10日胜率', '10日平均'], 1):
            cell = ws2.cell(row=row, column=col, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.border = border
        
        for stat in results['month_stats']:
            row += 1
            for col, h in enumerate(['月份', '信号数', '10日胜率', '10日平均'], 1):
                cell = ws2.cell(row=row, column=col, value=stat.get(h, ''))
                cell.border = border
    
    row += 3
    ws2.cell(row=row, column=1, value="【板块统计】").font = Font(bold=True, size=12)
    row += 2
    
    if results and 'sector_stats' in results:
        for col, h in enumerate(['板块', '信号数', '10日胜率', '10日平均'], 1):
            cell = ws2.cell(row=row, column=col, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.border = border
        
        for stat in results['sector_stats']:
            row += 1
            for col, h in enumerate(['板块', '信号数', '10日胜率', '10日平均'], 1):
                cell = ws2.cell(row=row, column=col, value=stat.get(h, ''))
                cell.border = border
    
    # 调整列宽
    ws1.column_dimensions['A'].width = 10
    ws1.column_dimensions['B'].width = 15
    ws1.column_dimensions['C'].width = 10
    ws1.column_dimensions['D'].width = 12
    ws1.column_dimensions['E'].width = 12
    for col in ['F', 'G', 'H', 'I', 'J', 'K', 'L', 'M']:
        ws1.column_dimensions[col].width = 10
    
    for col in range(1, 6):
        ws2.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 14
    
    wb.save(filename)
    print(f"\n✓ 已导出: {filename}")

def print_conclusion(results):
    """打印结论"""
    print("\n" + "=" * 70)
    print("【策略有效性结论】")
    print("=" * 70)
    
    if not results or 'period_stats' not in results:
        print("数据不足")
        return
    
    stats_10 = None
    for s in results['period_stats']:
        if s['持有期'] == '10日':
            stats_10 = s
            break
    
    if stats_10:
        win_rate = float(stats_10['胜率'].replace('%', ''))
        avg_ret = float(stats_10['平均收益'].replace('%', ''))
        
        print(f"""
核心指标（10日持有期）：
  - 胜率: {stats_10['胜率']}
  - 平均收益: {stats_10['平均收益']}
  - 样本数: {stats_10['样本数']}

策略评价：""")
        
        if win_rate >= 55 and avg_ret >= 2:
            print("  ✓ 策略有效！胜率和收益都不错")
        elif win_rate >= 50 or avg_ret >= 0:
            print("  ~ 策略一般，有一定参考价值")
        else:
            print("  ✗ 策略效果不佳，需要优化")

if __name__ == "__main__":
    signals = backtest_sector_washout_parallel(max_workers=10)
    results = analyze_results(signals)
    print_all_cases(signals)
    export_to_excel(signals, results)
    print_conclusion(results)
