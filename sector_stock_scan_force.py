"""
强制扫描版本 - 忽略市场情绪限制
用于在任何市场环境下寻找潜力板块和个股
"""
import akshare as ak
import pandas as pd
import numpy as np
from datetime import datetime, timedelta
from concurrent.futures import ThreadPoolExecutor, as_completed
import threading
import warnings
warnings.filterwarnings('ignore')

print_lock = threading.Lock()

def safe_print(msg):
    with print_lock:
        print(msg)

# ==================== 市场情绪分析 ====================
def analyze_market_sentiment():
    """详细分析当前市场情绪"""
    try:
        df = ak.stock_zh_index_daily(symbol="sh000001")
        df['date'] = pd.to_datetime(df['date'])
        df = df.tail(60).reset_index(drop=True)
        df['pct_change'] = df['close'].pct_change() * 100
        
        latest = df.iloc[-1]
        recent_5 = df.tail(5)
        recent_10 = df.tail(10)
        recent_20 = df.tail(20)
        
        # 计算各项指标
        ma5 = df['close'].tail(5).mean()
        ma10 = df['close'].tail(10).mean()
        ma20 = df['close'].tail(20).mean()
        
        vol_5d = recent_5['pct_change'].std()
        vol_10d = recent_10['pct_change'].std()
        vol_20d = recent_20['pct_change'].std()
        
        up_days_5 = (recent_5['pct_change'] > 0).sum()
        up_days_10 = (recent_10['pct_change'] > 0).sum()
        
        gain_5d = (latest['close'] / recent_5.iloc[0]['close'] - 1) * 100
        gain_10d = (latest['close'] / recent_10.iloc[0]['close'] - 1) * 100
        gain_20d = (latest['close'] / recent_20.iloc[0]['close'] - 1) * 100
        
        print("=" * 70)
        print("当前市场情绪分析")
        print("=" * 70)
        print(f"\n上证指数: {latest['close']:.2f}")
        print(f"\n均线系统:")
        print(f"  MA5:  {ma5:.2f} {'✓' if latest['close'] > ma5 else '✗'}")
        print(f"  MA10: {ma10:.2f} {'✓' if latest['close'] > ma10 else '✗'}")
        print(f"  MA20: {ma20:.2f} {'✓' if latest['close'] > ma20 else '✗'}")
        
        print(f"\n波动率:")
        print(f"  5日:  {vol_5d:.2f}% {'✓高' if vol_5d > 1.0 else '✗低'}")
        print(f"  10日: {vol_10d:.2f}%")
        print(f"  20日: {vol_20d:.2f}%")
        
        print(f"\n涨跌天数:")
        print(f"  近5日:  {up_days_5}/5天上涨")
        print(f"  近10日: {up_days_10}/10天上涨")
        
        print(f"\n累计涨幅:")
        print(f"  5日:  {gain_5d:+.2f}%")
        print(f"  10日: {gain_10d:+.2f}%")
        print(f"  20日: {gain_20d:+.2f}%")
        
        # 综合评分
        score = 0
        reasons = []
        
        if vol_5d >= 1.0:
            score += 2
            reasons.append("波动率适中")
        elif vol_5d < 0.5:
            reasons.append("波动率过低")
        
        if latest['close'] > ma5 > ma10:
            score += 2
            reasons.append("均线多头")
        elif latest['close'] < ma20:
            reasons.append("跌破MA20")
        
        if up_days_5 >= 3:
            score += 1
            reasons.append("近期多涨")
        
        if gain_10d > 2:
            score += 1
            reasons.append("10日上涨")
        
        print(f"\n综合评分: {score}/6")
        print(f"评价: {', '.join(reasons) if reasons else '无明显特征'}")
        
        if score >= 4:
            print("\n✓ 市场情绪良好，适合操作")
        elif score >= 2:
            print("\n⚠️ 市场情绪一般，谨慎操作")
        else:
            print("\n✗ 市场情绪较差，建议观望")
        
        return score, vol_5d
        
    except Exception as e:
        print(f"市场情绪分析失败: {e}")
        return 0, 0


# ==================== 板块扫描 ====================
def get_sector_history(sector_name, start_date, end_date):
    """获取板块历史数据"""
    try:
        df = ak.stock_board_concept_cons_em(symbol=sector_name)
        if df is None or len(df) == 0:
            return None
        
        top_stocks = df.head(5)['代码'].tolist()
        all_data = []
        
        for symbol in top_stocks:
            try:
                hist = ak.stock_zh_a_hist(
                    symbol=symbol,
                    period="daily",
                    start_date=start_date,
                    end_date=end_date,
                    adjust="qfq"
                )
                if hist is not None and len(hist) > 30:
                    hist['涨跌幅'] = hist['收盘'].pct_change() * 100
                    hist['日期'] = pd.to_datetime(hist['日期'])
                    all_data.append(hist[['日期', '涨跌幅']].set_index('日期'))
            except:
                continue
        
        if len(all_data) == 0:
            return None
        
        combined = pd.concat([d['涨跌幅'] for d in all_data], axis=1)
        return combined.mean(axis=1)
    except:
        return None


def detect_potential_sector(daily_returns):
    """检测潜力板块（宽松版）"""
    if daily_returns is None or len(daily_returns) < 60:
        return False, None
    
    daily_returns = daily_returns.dropna()
    if len(daily_returns) < 60:
        return False, None
    
    old_period = daily_returns.iloc[:-15]
    recent_15d = daily_returns.iloc[-15:]
    recent_5d = daily_returns.iloc[-5:]
    
    # 找爆发（降低标准到2.5%）
    surge_days = old_period[old_period > 2.5]
    if len(surge_days) == 0:
        return False, None
    
    max_surge_value = surge_days.max()
    max_surge_idx = surge_days.idxmax()
    surge_count = len(surge_days)
    
    after_surge = daily_returns[daily_returns.index >= max_surge_idx]
    if len(after_surge) < 10:
        return False, None
    
    # 计算回撤（放宽到5-20%）
    cumsum = after_surge.cumsum()
    max_gain = cumsum.max()
    current_gain = cumsum.iloc[-1]
    drawdown = max_gain - current_gain
    
    if drawdown < 5 or drawdown > 20:
        return False, None
    
    # 近期弱但企稳（放宽标准）
    recent_15d_sum = recent_15d.sum()
    recent_5d_sum = recent_5d.sum()
    recent_weak = recent_15d_sum < 8  # 放宽到8%
    not_crashing = recent_5d_sum > -8  # 放宽到-8%
    
    if not (recent_weak and not_crashing):
        return False, None
    
    return True, {
        'surge_date': max_surge_idx.strftime('%Y-%m-%d'),
        'surge_value': round(max_surge_value, 2),
        'surge_count': surge_count,
        'drawdown': round(drawdown, 2),
        'recent_15d': round(recent_15d_sum, 2),
        'recent_5d': round(recent_5d_sum, 2)
    }


# ==================== 个股扫描 ====================
def detect_potential_stock(df):
    """检测潜力个股（宽松版）"""
    if len(df) < 60:
        return False, None
    
    df = df.copy().reset_index(drop=True)
    df['涨跌幅'] = df['收盘'].pct_change() * 100
    
    latest_idx = len(df) - 1
    latest = df.iloc[-1]
    
    # 找大涨日（降低标准）
    surge_found = False
    surge_info = None
    
    for i in range(latest_idx - 15, max(latest_idx - 60, 0), -1):
        # 单日大涨>4%（降低标准）
        if df.iloc[i]['涨跌幅'] > 4:
            surge_found = True
            surge_info = {
                'idx': i,
                'date': df.iloc[i]['日期'],
                'gain': df.iloc[i]['涨跌幅'],
                'price': df.iloc[i]['收盘']
            }
            break
        
        # 或3日累计>8%（降低标准）
        if i >= 2:
            gain_3d = ((df.iloc[i]['收盘'] / df.iloc[i-3]['收盘']) - 1) * 100
            if gain_3d > 8:
                surge_found = True
                surge_info = {
                    'idx': i,
                    'date': df.iloc[i]['日期'],
                    'gain': gain_3d,
                    'price': df.iloc[i]['收盘']
                }
                break
    
    if not surge_found:
        return False, None
    
    # 计算回撤（放宽到6-25%）
    surge_idx = surge_info['idx']
    after_surge = df.iloc[surge_idx:]
    high_after = after_surge['最高'].max()
    current_price = latest['收盘']
    drawdown = (current_price - high_after) / high_after * 100
    
    if drawdown > -6 or drawdown < -25:
        return False, None
    
    # 近10日横盘（放宽到15%）
    recent_10 = df.tail(10)
    high_10 = recent_10['最高'].max()
    low_10 = recent_10['最低'].min()
    range_10 = (high_10 - low_10) / low_10 * 100
    
    if range_10 > 15:
        return False, None
    
    # 当前价格接近下沿（放宽到8%）
    dist_to_low = (current_price - low_10) / low_10 * 100
    if dist_to_low > 8:
        return False, None
    
    # 计算潜在收益空间
    potential_gain = (high_10 - current_price) / current_price * 100
    
    return True, {
        'surge_date': surge_info['date'],
        'surge_gain': round(surge_info['gain'], 2),
        'drawdown': round(drawdown, 2),
        'range_10d': round(range_10, 2),
        'dist_to_low': round(dist_to_low, 2),
        'low_10d': round(low_10, 2),
        'high_10d': round(high_10, 2),
        'potential_gain': round(potential_gain, 2)
    }


# ==================== 主扫描函数 ====================
def scan_market():
    """扫描市场"""
    # 分析市场情绪
    sentiment_score, volatility = analyze_market_sentiment()
    
    print("\n" + "=" * 70)
    print("开始扫描潜力板块")
    print("=" * 70)
    
    # 获取板块列表
    sectors_df = ak.stock_board_concept_name_em()
    sectors = sectors_df['板块名称'].tolist()[:100]
    
    potential_sectors = []
    
    for i, sector in enumerate(sectors):
        if (i + 1) % 20 == 0:
            safe_print(f"  进度: {i+1}/{len(sectors)}")
        
        data = get_sector_history(sector, '20250901', '20260115')
        is_potential, details = detect_potential_sector(data)
        
        if is_potential:
            potential_sectors.append({
                'name': sector,
                'details': details
            })
            safe_print(f"  ✓ {sector} - 回撤{details['drawdown']:.1f}%")
    
    print(f"\n发现 {len(potential_sectors)} 个潜力板块")
    
    if len(potential_sectors) == 0:
        print("未发现潜力板块")
        return []
    
    # 显示板块详情
    print("\n" + "=" * 70)
    print("潜力板块详情")
    print("=" * 70)
    for s in sorted(potential_sectors, key=lambda x: x['details']['drawdown'], reverse=True):
        d = s['details']
        print(f"\n{s['name']}:")
        print(f"  爆发日期: {d['surge_date']}, 涨幅{d['surge_value']}%")
        print(f"  回撤: {d['drawdown']}%")
        print(f"  近15日: {d['recent_15d']:+.1f}%, 近5日: {d['recent_5d']:+.1f}%")
    
    # 在潜力板块中扫描个股
    print("\n" + "=" * 70)
    print("在潜力板块中扫描个股")
    print("=" * 70)
    
    all_results = []
    completed = 0
    total_sectors = min(len(potential_sectors), 15)
    
    def scan_sector(sector_info):
        nonlocal completed
        sector = sector_info['name']
        sector_results = []
        
        try:
            cons = ak.stock_board_concept_cons_em(symbol=sector)
            if cons is None:
                return sector_results
            
            stocks = cons[cons['代码'].str.match(r'^(60|00)')].head(30)
            
            for _, stock in stocks.iterrows():
                code = stock['代码']
                name = stock['名称']
                
                try:
                    hist = ak.stock_zh_a_hist(
                        symbol=code,
                        period="daily",
                        start_date='20250901',
                        adjust="qfq"
                    )
                    
                    if hist is None or len(hist) < 60:
                        continue
                    
                    hist['日期'] = pd.to_datetime(hist['日期'])
                    is_potential, stock_details = detect_potential_stock(hist)
                    
                    if is_potential:
                        sector_results.append({
                            '板块': sector,
                            '代码': code,
                            '名称': name,
                            '现价': hist.iloc[-1]['收盘'],
                            '今日涨幅': stock.get('涨跌幅', 0),
                            **stock_details
                        })
                        safe_print(f"  ✓ {sector} - {code} {name} 距低点{stock_details['dist_to_low']:.1f}%")
                except:
                    pass
        except:
            pass
        
        with print_lock:
            completed += 1
            if completed % 3 == 0:
                safe_print(f"\n  板块进度: {completed}/{total_sectors}\n")
        
        return sector_results
    
    with ThreadPoolExecutor(max_workers=5) as executor:
        futures = [executor.submit(scan_sector, s) for s in potential_sectors[:total_sectors]]
        for future in as_completed(futures):
            all_results.extend(future.result())
    
    return all_results, sentiment_score, volatility


# ==================== 导出结果 ====================
def export_results(results, sentiment_score, volatility):
    """导出结果到Excel"""
    if len(results) == 0:
        print("\n未发现符合条件的个股")
        return
    
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    
    results_df = pd.DataFrame(results)
    
    # 按潜在收益排序
    results_df = results_df.sort_values('potential_gain', ascending=False)
    
    filename = f"板块个股扫描_强制_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "扫描结果"
    
    # 样式
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    green_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    yellow_fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
    center = Alignment(horizontal="center", vertical="center")
    border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )
    
    # 表头
    headers = ['排名', '板块', '代码', '名称', '现价', '今日涨幅', 
               '爆发日期', '爆发涨幅', '回撤', '距低点', '潜在收益']
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center
        cell.border = border
    
    # 数据
    for row_idx, (_, r) in enumerate(results_df.iterrows(), 2):
        row_data = [
            row_idx - 1,
            r['板块'],
            r['代码'],
            r['名称'],
            round(r['现价'], 2),
            f"{r['今日涨幅']:.2f}%",
            r['surge_date'],
            f"{r['surge_gain']:.1f}%",
            f"{r['drawdown']:.1f}%",
            f"{r['dist_to_low']:.1f}%",
            f"{r['potential_gain']:.1f}%"
        ]
        
        for col, val in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col, value=val)
            cell.alignment = center
            cell.border = border
            
            # 根据潜在收益着色
            if col == 11:  # 潜在收益列
                if r['potential_gain'] > 8:
                    cell.fill = green_fill
                elif r['potential_gain'] > 5:
                    cell.fill = yellow_fill
    
    # 调整列宽
    ws.column_dimensions['A'].width = 6
    ws.column_dimensions['B'].width = 14
    ws.column_dimensions['C'].width = 10
    ws.column_dimensions['D'].width = 12
    ws.column_dimensions['E'].width = 10
    ws.column_dimensions['F'].width = 10
    ws.column_dimensions['G'].width = 12
    ws.column_dimensions['H'].width = 10
    ws.column_dimensions['I'].width = 10
    ws.column_dimensions['J'].width = 10
    ws.column_dimensions['K'].width = 10
    
    # 添加市场情绪说明
    ws2 = wb.create_sheet(title="市场情绪")
    ws2['A1'] = "市场情绪分析"
    ws2['A1'].font = Font(bold=True, size=14)
    ws2['A3'] = "情绪评分"
    ws2['B3'] = f"{sentiment_score}/6"
    ws2['A4'] = "波动率"
    ws2['B4'] = f"{volatility:.2f}%"
    ws2['A6'] = "说明"
    ws2['A7'] = "评分>=4: 市场情绪良好"
    ws2['A8'] = "评分2-3: 市场情绪一般"
    ws2['A9'] = "评分<2: 市场情绪较差"
    
    wb.save(filename)
    print(f"\n✓ 结果已保存: {filename}")
    
    # 打印摘要
    print("\n" + "=" * 70)
    print("扫描摘要")
    print("=" * 70)
    print(f"\n发现 {len(results)} 只潜力个股")
    print(f"市场情绪评分: {sentiment_score}/6")
    print(f"市场波动率: {volatility:.2f}%")
    
    # 按板块统计
    print("\n按板块统计:")
    sector_counts = results_df['板块'].value_counts()
    for sector, count in sector_counts.head(10).items():
        print(f"  {sector}: {count}只")
    
    # Top 10
    print("\nTop 10 潜在收益:")
    for i, (_, r) in enumerate(results_df.head(10).iterrows(), 1):
        print(f"  {i}. {r['代码']} {r['名称']} [{r['板块']}] "
              f"潜在收益{r['potential_gain']:.1f}% 距低点{r['dist_to_low']:.1f}%")


if __name__ == "__main__":
    results, sentiment_score, volatility = scan_market()
    export_results(results, sentiment_score, volatility)
