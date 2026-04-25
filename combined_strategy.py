"""
组合策略：潜力板块 + 涨停破位洗盘
逻辑：在潜力板块中找涨停破位洗盘反包的个股，双重筛选提高胜率
"""
import akshare as ak
import pandas as pd
import numpy as np
from datetime import datetime
from strategies import strategy_limit_up_washout

def get_potential_sectors():
    """获取当前潜力板块"""
    return ['GDR', '电子后视镜', '举牌', '乳业', '3D摄像头', '低碳冶金']

def get_sector_stocks(sector_name):
    """获取板块成分股"""
    try:
        df = ak.stock_board_concept_cons_em(symbol=sector_name)
        if df is None:
            return []
        df = df[df['代码'].str.match(r'^(60|00)')]
        df = df[~df['名称'].str.contains('ST')]
        return df[['代码', '名称', '涨跌幅']].to_dict('records')
    except:
        return []

def check_limit_up_washout(symbol, name, relaxed=True):
    """
    检查个股是否符合涨停破位洗盘战法
    使用strategies.py中的strategy_limit_up_washout函数
    """
    try:
        df = ak.stock_zh_a_hist(symbol=symbol, period="daily", adjust="qfq")
        if df is None or len(df) < 30:
            return None, None
        
        df = df.tail(70)
        result = strategy_limit_up_washout(df, relaxed=relaxed)
        
        if result['触发']:
            return {
                '代码': symbol,
                '名称': name,
                '说明': result['说明'],
                '条件': result['条件'],
                '状态': '完全匹配'
            }, None
        
        # 检查类似形态（接近触发但还没完全触发）
        similar = check_similar_pattern(df, symbol, name, relaxed)
        return None, similar
        
    except Exception as e:
        return None, None

def check_similar_pattern(df, symbol, name, relaxed=True):
    """检查类似形态：涨停后回调，接近破位或已破位但还没反包"""
    if len(df) < 20:
        return None
    
    df = df.copy().reset_index(drop=True)
    latest = df.iloc[-1]
    latest_idx = len(df) - 1
    
    search_days = 30 if relaxed else 20
    
    # 找涨停板
    limit_up_info = None
    for i in range(latest_idx - 3, max(latest_idx - search_days, 0), -1):
        row = df.iloc[i]
        prev_close = df.iloc[i-1]['收盘'] if i > 0 else row['开盘']
        gain = (row['收盘'] - prev_close) / prev_close * 100
        
        if gain >= 9.5 and row['开盘'] < row['收盘'] * 0.99:
            limit_up_info = {
                'idx': i,
                'date': row.get('日期', f'第{i}天'),
                'low': row['最低'],
                'close': row['收盘']
            }
            break
    
    if not limit_up_info:
        return None
    
    # 检查是否先创新高（排除）
    for i in range(limit_up_info['idx'] + 1, latest_idx + 1):
        if df.iloc[i]['最高'] > limit_up_info['close'] * 1.08:
            return None
    
    limit_low = limit_up_info['low']
    
    # 检查是否破位
    broke_low = False
    for i in range(limit_up_info['idx'] + 1, latest_idx + 1):
        if df.iloc[i]['最低'] < limit_low:
            broke_low = True
            break
    
    # 当前位置
    distance_to_low = (latest['收盘'] - limit_low) / limit_low * 100
    today_gain = (latest['收盘'] - latest['开盘']) / latest['开盘'] * 100
    
    # 类似形态条件
    if broke_low and -10 < distance_to_low < 5:
        # 已破位，等待反包
        return {
            '代码': symbol,
            '名称': name,
            '说明': f"涨停日:{limit_up_info['date']}, 已破位待反包",
            '条件': f"涨停最低价{limit_low:.2f}, 当前距离{distance_to_low:+.1f}%, 今日{today_gain:+.1f}%",
            '状态': '待反包',
            '涨停最低价': limit_low,
            '距离': distance_to_low
        }
    elif not broke_low and 0 < distance_to_low < 8:
        # 接近破位
        return {
            '代码': symbol,
            '名称': name,
            '说明': f"涨停日:{limit_up_info['date']}, 接近破位",
            '条件': f"涨停最低价{limit_low:.2f}, 当前距离{distance_to_low:+.1f}%, 今日{today_gain:+.1f}%",
            '状态': '接近破位',
            '涨停最低价': limit_low,
            '距离': distance_to_low
        }
    
    return None

def scan_combined_strategy(relaxed=True):
    """扫描组合策略：潜力板块 + 涨停破位洗盘"""
    print("=" * 70)
    print("组合策略扫描：潜力板块 + 涨停破位洗盘")
    print("=" * 70)
    print("\n策略逻辑：")
    print("  第一层：筛选潜力板块（早期爆发→回撤洗盘→企稳）")
    print("  第二层：在板块中找涨停破位洗盘反包的个股")
    print("  双重共振 = 板块要启动 + 个股有买点")
    
    mode = "宽松模式" if relaxed else "严格模式"
    print(f"\n扫描模式: {mode}")
    
    sectors = get_potential_sectors()
    print(f"潜力板块: {', '.join(sectors)}")
    
    all_matches = []  # 完全匹配
    all_similar = []  # 类似形态
    
    for sector in sectors:
        print(f"\n{'='*50}")
        print(f"【{sector}】")
        print("=" * 50)
        
        stocks = get_sector_stocks(sector)
        if not stocks:
            print("  获取成分股失败")
            continue
        
        print(f"  成分股数量: {len(stocks)}")
        
        for stock in stocks:
            code = stock['代码']
            name = stock['名称']
            
            match, similar = check_limit_up_washout(code, name, relaxed)
            
            if match:
                match['板块'] = sector
                match['今日涨幅'] = stock.get('涨跌幅', 0)
                all_matches.append(match)
                print(f"  ✓ 完全匹配: {code} {name}")
            elif similar:
                similar['板块'] = sector
                similar['今日涨幅'] = stock.get('涨跌幅', 0)
                all_similar.append(similar)
                print(f"  ~ 类似形态: {code} {name} ({similar['状态']})")
    
    return all_matches, all_similar

def print_results(matches, similar):
    """打印结果"""
    print("\n" + "=" * 70)
    print("扫描结果")
    print("=" * 70)
    
    # 完全匹配
    print(f"\n【完全匹配】{len(matches)}只 - 今日可买入")
    print("-" * 60)
    if matches:
        for m in matches:
            print(f"\n  {m['代码']} {m['名称']} [{m['板块']}]")
            print(f"    {m['说明']}")
            print(f"    {m['条件']}")
    else:
        print("  暂无完全匹配的个股")
    
    # 类似形态
    print(f"\n【类似形态】{len(similar)}只 - 待观察")
    print("-" * 60)
    if similar:
        # 按状态分组
        waiting = [s for s in similar if s['状态'] == '待反包']
        approaching = [s for s in similar if s['状态'] == '接近破位']
        
        if waiting:
            print(f"\n  已破位待反包 ({len(waiting)}只):")
            for s in sorted(waiting, key=lambda x: x['距离']):
                print(f"    {s['代码']} {s['名称']} [{s['板块']}] 距离涨停最低价{s['距离']:+.1f}%")
        
        if approaching:
            print(f"\n  接近破位 ({len(approaching)}只):")
            for s in sorted(approaching, key=lambda x: x['距离']):
                print(f"    {s['代码']} {s['名称']} [{s['板块']}] 距离涨停最低价{s['距离']:+.1f}%")
    else:
        print("  暂无类似形态的个股")
    
    # 策略说明
    print("\n" + "=" * 70)
    print("【组合策略说明】")
    print("=" * 70)
    print("""
双重筛选逻辑：
  板块层面：潜力板块（早期爆发→回撤12-16%→企稳）
  个股层面：涨停破位洗盘（涨停→跌破最低价→大阳反包）

买入条件：
  1. 完全匹配：今日大阳反包，可以买入
  2. 待反包：已破位，等大阳线出现再买
  3. 接近破位：还没破位，继续观察

止损止盈：
  止损：跌破涨停最低价5%
  止盈：反弹到涨停收盘价（10-20%空间）
""")

def export_to_excel(matches, similar):
    """导出到Excel"""
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    
    filename = f"组合策略_潜力板块+涨停洗盘_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    
    wb = openpyxl.Workbook()
    
    # 样式
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    green_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    yellow_fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
    center = Alignment(horizontal="center", vertical="center")
    border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )
    
    # Sheet1: 完全匹配
    ws1 = wb.active
    ws1.title = "完全匹配-可买入"
    
    headers = ["板块", "代码", "名称", "今日涨幅", "触发说明", "详细条件"]
    for col, h in enumerate(headers, 1):
        cell = ws1.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center
        cell.border = border
    
    for row_idx, m in enumerate(matches, 2):
        row_data = [m['板块'], m['代码'], m['名称'], 
                    f"{m.get('今日涨幅', 0):.2f}%", m['说明'], m['条件']]
        for col, val in enumerate(row_data, 1):
            cell = ws1.cell(row=row_idx, column=col, value=val)
            cell.alignment = center if col <= 4 else Alignment(horizontal="left")
            cell.border = border
            cell.fill = green_fill
    
    # Sheet2: 类似形态
    ws2 = wb.create_sheet(title="类似形态-待观察")
    
    headers2 = ["板块", "代码", "名称", "状态", "今日涨幅", "涨停最低价", "距离", "说明"]
    for col, h in enumerate(headers2, 1):
        cell = ws2.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center
        cell.border = border
    
    for row_idx, s in enumerate(similar, 2):
        row_data = [s['板块'], s['代码'], s['名称'], s['状态'],
                    f"{s.get('今日涨幅', 0):.2f}%", 
                    s.get('涨停最低价', ''),
                    f"{s.get('距离', 0):+.1f}%",
                    s['说明']]
        for col, val in enumerate(row_data, 1):
            cell = ws2.cell(row=row_idx, column=col, value=val)
            cell.alignment = center
            cell.border = border
            if s['状态'] == '待反包':
                cell.fill = yellow_fill
    
    # 调整列宽
    for ws in [ws1, ws2]:
        ws.column_dimensions['A'].width = 12
        ws.column_dimensions['B'].width = 10
        ws.column_dimensions['C'].width = 12
        ws.column_dimensions['D'].width = 10
        ws.column_dimensions['E'].width = 40
        ws.column_dimensions['F'].width = 50
        if ws == ws2:
            ws.column_dimensions['G'].width = 10
            ws.column_dimensions['H'].width = 30
    
    wb.save(filename)
    print(f"\n✓ 已导出到: {filename}")

if __name__ == "__main__":
    matches, similar = scan_combined_strategy(relaxed=True)
    print_results(matches, similar)
    export_to_excel(matches, similar)
