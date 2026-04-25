"""
MACD金叉策略扫描
检测MACD线上穿信号线（DEA）的金叉信号
"""
import akshare as ak
import pandas as pd
import numpy as np
from datetime import datetime, timedelta
from concurrent.futures import ThreadPoolExecutor, as_completed
import threading
import warnings
warnings.filterwarnings('ignore')

print_lock = threading.Lock()

def safe_print(msg):
    with print_lock:
        print(msg)


def calculate_macd(df, fast=12, slow=26, signal=9):
    """
    计算MACD指标
    
    参数:
        fast: 快线周期（默认12）
        slow: 慢线周期（默认26）
        signal: 信号线周期（默认9）
    
    返回:
        df: 包含MACD、Signal、Histogram的DataFrame
    """
    df = df.copy()
    
    # 计算EMA
    ema_fast = df['收盘'].ewm(span=fast, adjust=False).mean()
    ema_slow = df['收盘'].ewm(span=slow, adjust=False).mean()
    
    # MACD线 = 快线 - 慢线
    df['MACD'] = ema_fast - ema_slow
    
    # Signal线（DEA）= MACD的9日EMA
    df['Signal'] = df['MACD'].ewm(span=signal, adjust=False).mean()
    
    # 柱状图（MACD柱）= MACD - Signal
    df['Histogram'] = df['MACD'] - df['Signal']
    
    return df


def detect_macd_golden_cross(df, strict=True):
    """
    检测MACD金叉信号
    
    参数:
        df: 股票数据
        strict: 是否严格模式
            严格模式：MACD和Signal都在0轴上方
            宽松模式：允许在0轴下方金叉
    
    返回:
        is_golden_cross: 是否金叉
        details: 详细信息
    """
    if len(df) < 50:
        return False, None
    
    df = calculate_macd(df)
    
    if len(df) < 2:
        return False, None
    
    latest = df.iloc[-1]
    prev = df.iloc[-2]
    
    # 检查是否有NaN
    if pd.isna(latest['MACD']) or pd.isna(latest['Signal']):
        return False, None
    
    # 金叉条件：MACD上穿Signal
    is_cross = (latest['MACD'] > latest['Signal']) and (prev['MACD'] <= prev['Signal'])
    
    if not is_cross:
        return False, None
    
    # 严格模式：要求在0轴上方
    if strict:
        if latest['MACD'] < 0 or latest['Signal'] < 0:
            return False, None
    
    # 计算金叉强度
    cross_strength = latest['MACD'] - latest['Signal']
    
    # 检查MACD趋势（是否向上）
    macd_trend = latest['MACD'] - prev['MACD']
    
    # 检查柱状图（红柱）
    histogram_positive = latest['Histogram'] > 0
    
    # 计算距离0轴的距离
    dist_to_zero = latest['MACD']
    
    # 检查是否是底部金叉（MACD在负值区域）
    is_bottom_cross = latest['MACD'] < 0
    
    # 检查前期是否有死叉
    death_cross_idx = None
    for i in range(len(df) - 2, max(len(df) - 20, 0), -1):
        if df.iloc[i]['MACD'] < df.iloc[i]['Signal'] and df.iloc[i-1]['MACD'] >= df.iloc[i-1]['Signal']:
            death_cross_idx = i
            break
    
    # 计算金叉后的潜在空间
    recent_high = df['收盘'].tail(20).max()
    current_price = latest['收盘']
    potential_gain = (recent_high - current_price) / current_price * 100
    
    return True, {
        'MACD': round(latest['MACD'], 4),
        'Signal': round(latest['Signal'], 4),
        'Histogram': round(latest['Histogram'], 4),
        'cross_strength': round(cross_strength, 4),
        'macd_trend': round(macd_trend, 4),
        'histogram_positive': histogram_positive,
        'dist_to_zero': round(dist_to_zero, 4),
        'is_bottom_cross': is_bottom_cross,
        'death_cross_days_ago': len(df) - death_cross_idx - 1 if death_cross_idx else None,
        'potential_gain': round(potential_gain, 2),
        'position': '0轴上方' if latest['MACD'] > 0 else '0轴下方'
    }


def get_all_stocks():
    """获取所有A股主板股票"""
    try:
        df = ak.stock_zh_a_spot_em()
        # 主板股票
        df = df[df['代码'].str.match(r'^(60|00)')]
        # 排除ST
        df = df[~df['名称'].str.contains('ST')]
        # 有效数据
        df = df[df['最新价'].notna() & (df['最新价'] > 0)]
        df = df[df['成交量'].notna() & (df['成交量'] > 0)]
        return df[['代码', '名称', '最新价', '涨跌幅', '换手率']].to_dict('records')
    except Exception as e:
        print(f"获取股票列表失败: {e}")
        return []


def scan_stock(stock, strict=True):
    """扫描单只股票"""
    code = stock['代码']
    name = stock['名称']
    
    try:
        # 获取历史数据
        df = ak.stock_zh_a_hist(
            symbol=code,
            period="daily",
            start_date=(datetime.now() - timedelta(days=180)).strftime('%Y%m%d'),
            adjust="qfq"
        )
        
        if df is None or len(df) < 50:
            return None
        
        df['日期'] = pd.to_datetime(df['日期'])
        
        # 检测金叉
        is_golden, details = detect_macd_golden_cross(df, strict=strict)
        
        if is_golden:
            return {
                '代码': code,
                '名称': name,
                '现价': stock['最新价'],
                '今日涨幅': stock['涨跌幅'],
                '换手率': stock['换手率'],
                **details
            }
        
        return None
        
    except Exception as e:
        return None


def scan_market(strict=True, max_workers=10):
    """扫描全市场"""
    print("=" * 70)
    print("MACD金叉策略扫描")
    print("=" * 70)
    print(f"模式: {'严格模式（0轴上方）' if strict else '宽松模式（允许0轴下方）'}")
    
    print("\n获取股票列表...")
    stocks = get_all_stocks()
    print(f"共 {len(stocks)} 只股票")
    
    print(f"\n开始扫描（{max_workers}线程）...")
    print("=" * 70)
    
    results = []
    completed = 0
    lock = threading.Lock()
    
    def worker(stock):
        nonlocal completed
        result = scan_stock(stock, strict=strict)
        
        with lock:
            completed += 1
            if result:
                results.append(result)
                safe_print(f"[{completed}/{len(stocks)}] ✓ {result['代码']} {result['名称']} "
                          f"MACD:{result['MACD']:.4f} {result['position']}")
            elif completed % 200 == 0:
                safe_print(f"[{completed}/{len(stocks)}] 进度 {completed/len(stocks)*100:.1f}%")
        
        return result
    
    with ThreadPoolExecutor(max_workers=max_workers) as executor:
        futures = [executor.submit(worker, stock) for stock in stocks]
        for future in as_completed(futures):
            pass
    
    print("\n" + "=" * 70)
    print(f"扫描完成！发现 {len(results)} 个MACD金叉信号")
    print("=" * 70)
    
    return results


def print_results(results):
    """打印结果"""
    if not results:
        print("\n未发现MACD金叉信号")
        return
    
    # 按位置分组
    above_zero = [r for r in results if r['position'] == '0轴上方']
    below_zero = [r for r in results if r['position'] == '0轴下方']
    
    print(f"\n【0轴上方金叉】{len(above_zero)}只 - 强势上涨信号")
    print("-" * 70)
    if above_zero:
        # 按MACD值排序
        above_zero_sorted = sorted(above_zero, key=lambda x: x['MACD'], reverse=True)
        for r in above_zero_sorted[:20]:  # 显示前20只
            print(f"  {r['代码']} {r['名称']:<8} "
                  f"MACD:{r['MACD']:>7.4f} Signal:{r['Signal']:>7.4f} "
                  f"柱:{r['Histogram']:>6.4f} 涨幅:{r['今日涨幅']:>+6.2f}%")
    
    print(f"\n【0轴下方金叉】{len(below_zero)}只 - 底部反转信号")
    print("-" * 70)
    if below_zero:
        # 按距离0轴距离排序（越接近0轴越好）
        below_zero_sorted = sorted(below_zero, key=lambda x: abs(x['MACD']))
        for r in below_zero_sorted[:20]:  # 显示前20只
            print(f"  {r['代码']} {r['名称']:<8} "
                  f"MACD:{r['MACD']:>7.4f} Signal:{r['Signal']:>7.4f} "
                  f"柱:{r['Histogram']:>6.4f} 涨幅:{r['今日涨幅']:>+6.2f}%")
    
    # 统计分析
    print("\n" + "=" * 70)
    print("统计分析")
    print("=" * 70)
    
    avg_macd = np.mean([r['MACD'] for r in results])
    avg_gain_today = np.mean([r['今日涨幅'] for r in results])
    avg_potential = np.mean([r['potential_gain'] for r in results])
    
    print(f"\n平均MACD值: {avg_macd:.4f}")
    print(f"平均今日涨幅: {avg_gain_today:+.2f}%")
    print(f"平均潜在收益: {avg_potential:.2f}%")
    
    # 按潜在收益排序
    print(f"\n【Top 10 潜在收益】")
    print("-" * 70)
    top_potential = sorted(results, key=lambda x: x['potential_gain'], reverse=True)[:10]
    for i, r in enumerate(top_potential, 1):
        print(f"  {i}. {r['代码']} {r['名称']:<8} "
              f"潜在收益:{r['potential_gain']:>6.2f}% "
              f"MACD:{r['MACD']:>7.4f} {r['position']}")


def export_to_excel(results, strict=True):
    """导出到Excel"""
    if not results:
        print("没有数据可导出")
        return
    
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    
    results_df = pd.DataFrame(results)
    
    # 按MACD值排序
    results_df = results_df.sort_values('MACD', ascending=False)
    
    mode = "严格" if strict else "宽松"
    filename = f"MACD金叉扫描_{mode}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    
    wb = openpyxl.Workbook()
    
    # 样式
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    green_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    yellow_fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
    red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    center = Alignment(horizontal="center", vertical="center")
    border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )
    
    # Sheet1: 0轴上方
    ws1 = wb.active
    ws1.title = "0轴上方金叉"
    
    headers = ['排名', '代码', '名称', '现价', '今日涨幅', '换手率', 
               'MACD', 'Signal', '柱状图', '金叉强度', '潜在收益']
    for col, h in enumerate(headers, 1):
        cell = ws1.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center
        cell.border = border
    
    above_zero = results_df[results_df['position'] == '0轴上方']
    for row_idx, (_, r) in enumerate(above_zero.iterrows(), 2):
        row_data = [
            row_idx - 1,
            r['代码'],
            r['名称'],
            round(r['现价'], 2),
            f"{r['今日涨幅']:.2f}%",
            f"{r['换手率']:.2f}%",
            round(r['MACD'], 4),
            round(r['Signal'], 4),
            round(r['Histogram'], 4),
            round(r['cross_strength'], 4),
            f"{r['potential_gain']:.2f}%"
        ]
        
        for col, val in enumerate(row_data, 1):
            cell = ws1.cell(row=row_idx, column=col, value=val)
            cell.alignment = center
            cell.border = border
            
            # 着色
            if col == 5:  # 今日涨幅
                if r['今日涨幅'] > 3:
                    cell.fill = green_fill
                elif r['今日涨幅'] < -3:
                    cell.fill = red_fill
            elif col == 7:  # MACD
                if r['MACD'] > 0.5:
                    cell.fill = green_fill
                elif r['MACD'] > 0.1:
                    cell.fill = yellow_fill
    
    # Sheet2: 0轴下方
    ws2 = wb.create_sheet(title="0轴下方金叉")
    
    for col, h in enumerate(headers, 1):
        cell = ws2.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center
        cell.border = border
    
    below_zero = results_df[results_df['position'] == '0轴下方']
    below_zero = below_zero.sort_values('MACD', ascending=False)  # 按MACD排序（越接近0越好）
    
    for row_idx, (_, r) in enumerate(below_zero.iterrows(), 2):
        row_data = [
            row_idx - 1,
            r['代码'],
            r['名称'],
            round(r['现价'], 2),
            f"{r['今日涨幅']:.2f}%",
            f"{r['换手率']:.2f}%",
            round(r['MACD'], 4),
            round(r['Signal'], 4),
            round(r['Histogram'], 4),
            round(r['cross_strength'], 4),
            f"{r['potential_gain']:.2f}%"
        ]
        
        for col, val in enumerate(row_data, 1):
            cell = ws2.cell(row=row_idx, column=col, value=val)
            cell.alignment = center
            cell.border = border
            
            if col == 5:  # 今日涨幅
                if r['今日涨幅'] > 3:
                    cell.fill = green_fill
                elif r['今日涨幅'] < -3:
                    cell.fill = red_fill
    
    # 调整列宽
    for ws in [ws1, ws2]:
        ws.column_dimensions['A'].width = 6
        ws.column_dimensions['B'].width = 10
        ws.column_dimensions['C'].width = 12
        ws.column_dimensions['D'].width = 10
        ws.column_dimensions['E'].width = 10
        ws.column_dimensions['F'].width = 10
        ws.column_dimensions['G'].width = 10
        ws.column_dimensions['H'].width = 10
        ws.column_dimensions['I'].width = 10
        ws.column_dimensions['J'].width = 10
        ws.column_dimensions['K'].width = 10
    
    # Sheet3: 统计分析
    ws3 = wb.create_sheet(title="统计分析")
    
    stats = [
        ['指标', '数值'],
        ['总信号数', len(results_df)],
        ['0轴上方', len(above_zero)],
        ['0轴下方', len(below_zero)],
        ['平均MACD', round(results_df['MACD'].mean(), 4)],
        ['平均今日涨幅', f"{results_df['今日涨幅'].mean():.2f}%"],
        ['平均潜在收益', f"{results_df['potential_gain'].mean():.2f}%"],
        ['最大MACD', round(results_df['MACD'].max(), 4)],
        ['最小MACD', round(results_df['MACD'].min(), 4)],
    ]
    
    for row_idx, row_data in enumerate(stats, 1):
        for col, val in enumerate(row_data, 1):
            cell = ws3.cell(row=row_idx, column=col, value=val)
            if row_idx == 1:
                cell.font = header_font
                cell.fill = header_fill
            cell.alignment = center
            cell.border = border
    
    ws3.column_dimensions['A'].width = 15
    ws3.column_dimensions['B'].width = 15
    
    wb.save(filename)
    print(f"\n✓ 结果已保存: {filename}")


def main():
    import argparse
    
    parser = argparse.ArgumentParser(description='MACD金叉策略扫描')
    parser.add_argument('--mode', choices=['strict', 'loose'], default='loose',
                       help='扫描模式: strict=严格(0轴上方), loose=宽松(允许0轴下方)')
    parser.add_argument('--workers', type=int, default=10, help='线程数')
    args = parser.parse_args()
    
    strict = (args.mode == 'strict')
    
    results = scan_market(strict=strict, max_workers=args.workers)
    print_results(results)
    export_to_excel(results, strict=strict)


if __name__ == "__main__":
    main()
