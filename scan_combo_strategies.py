"""
策略组合扫描 - 日常使用
基于回测验证的最佳策略组合扫描全市场
"""
import akshare as ak
import pandas as pd
import numpy as np
from datetime import datetime
from concurrent.futures import ThreadPoolExecutor, as_completed
from threading import Lock
import time

print_lock = Lock()

def safe_print(msg):
    with print_lock:
        print(msg)

# ============ 5种基础策略 ============

def strategy_shrink_consolidation(df: pd.DataFrame) -> dict:
    """缩量横盘整理"""
    if len(df) < 35:
        return {"触发": False}
    
    df = df.copy().reset_index(drop=True)
    df['MA20'] = df['收盘'].rolling(20).mean()
    latest = df.iloc[-1]
    
    conditions = []
    
    price_30d_ago = df.iloc[-30]['收盘']
    price_10d_ago = df.iloc[-10]['收盘']
    prior_gain = (price_10d_ago - price_30d_ago) / price_30d_ago * 100
    c1 = prior_gain > 8
    conditions.append(c1)
    
    recent_10 = df.tail(10)
    high_10 = recent_10['最高'].max()
    low_10 = recent_10['最低'].min()
    range_10 = (high_10 - low_10) / low_10 * 100
    c2 = range_10 < 15
    conditions.append(c2)
    
    vol_recent_5 = df['成交量'].tail(5).mean()
    vol_prev_20 = df['成交量'].iloc[-25:-5].mean() if len(df) >= 25 else df['成交量'].mean()
    vol_ratio = vol_recent_5 / vol_prev_20 if vol_prev_20 > 0 else 1
    c3 = vol_ratio < 0.9
    conditions.append(c3)
    
    c4 = latest['收盘'] > latest['MA20'] if pd.notna(latest['MA20']) else False
    conditions.append(c4)
    
    return {"触发": sum(conditions) >= 3}


def strategy_deep_pullback_support(df: pd.DataFrame) -> dict:
    """深度回调支撑"""
    if len(df) < 35:
        return {"触发": False}
    
    df = df.copy().reset_index(drop=True)
    df['MA20'] = df['收盘'].rolling(20).mean()
    latest = df.iloc[-1]
    
    conditions = []
    
    high_30d = df.tail(30)['最高'].max()
    current_price = latest['收盘']
    pullback = (current_price - high_30d) / high_30d * 100
    c1 = pullback < -8
    conditions.append(c1)
    
    ma20 = latest['MA20']
    c2 = abs(current_price - ma20) / ma20 < 0.08 if pd.notna(ma20) else False
    conditions.append(c2)
    
    is_red = latest['收盘'] > latest['开盘']
    body = abs(latest['收盘'] - latest['开盘'])
    lower_shadow = min(latest['开盘'], latest['收盘']) - latest['最低']
    c3 = is_red or (lower_shadow > body * 0.3)
    conditions.append(c3)
    
    return {"触发": sum(conditions) >= 3}


def strategy_box_breakout(df: pd.DataFrame) -> dict:
    """箱体突破"""
    if len(df) < 20:
        return {"触发": False}
    
    df = df.copy().reset_index(drop=True)
    latest = df.iloc[-1]
    
    box_data = df.iloc[-16:-1]
    box_high = box_data['最高'].max()
    box_low = box_data['最低'].min()
    box_range = (box_high - box_low) / box_low * 100
    
    conditions = []
    c1 = 5 < box_range < 30
    conditions.append(c1)
    
    c2 = latest['收盘'] > box_high
    conditions.append(c2)
    
    vol_ma5 = df['成交量'].iloc[-6:-1].mean()
    vol_ratio = latest['成交量'] / vol_ma5 if vol_ma5 > 0 else 0
    c3 = vol_ratio > 1.0
    conditions.append(c3)
    
    return {"触发": sum(conditions) >= 3}


def strategy_ma_convergence_divergence(df: pd.DataFrame) -> dict:
    """均线粘合发散"""
    if len(df) < 25:
        return {"触发": False}
    
    df = df.copy().reset_index(drop=True)
    df['MA5'] = df['收盘'].rolling(5).mean()
    df['MA10'] = df['收盘'].rolling(10).mean()
    df['MA20'] = df['收盘'].rolling(20).mean()
    
    latest = df.iloc[-1]
    prev = df.iloc[-2]
    
    if pd.isna(prev['MA5']) or pd.isna(prev['MA10']) or pd.isna(prev['MA20']):
        return {"触发": False}
    
    conditions = []
    
    ma_values = [prev['MA5'], prev['MA10'], prev['MA20']]
    ma_max = max(ma_values)
    ma_min = min(ma_values)
    ma_spread = (ma_max - ma_min) / ma_min * 100
    c1 = ma_spread < 8
    conditions.append(c1)
    
    vol_ma5 = df['成交量'].iloc[-6:-1].mean()
    vol_ratio = latest['成交量'] / vol_ma5 if vol_ma5 > 0 else 0
    c2 = vol_ratio > 1.0
    conditions.append(c2)
    
    today_gain = (latest['收盘'] - latest['开盘']) / latest['开盘'] * 100
    c3 = today_gain > 0.5
    conditions.append(c3)
    
    return {"触发": sum(conditions) >= 3}


def strategy_volume_bottom(df: pd.DataFrame) -> dict:
    """地量见底"""
    if len(df) < 25:
        return {"触发": False}
    
    df = df.copy().reset_index(drop=True)
    latest = df.iloc[-1]
    
    conditions = []
    
    vol_recent_5 = df['成交量'].tail(6).iloc[:-1].mean()
    vol_prev_20 = df['成交量'].iloc[-25:-5].mean() if len(df) >= 25 else df['成交量'].mean()
    vol_ratio_5d = vol_recent_5 / vol_prev_20 if vol_prev_20 > 0 else 1
    c1 = vol_ratio_5d < 0.8
    conditions.append(c1)
    
    high_20d = df['最高'].tail(20).max()
    current_price = latest['收盘']
    pullback = (current_price - high_20d) / high_20d * 100
    c2 = pullback < -5
    conditions.append(c2)
    
    vol_ma5 = df['成交量'].iloc[-6:-1].mean()
    vol_ratio = latest['成交量'] / vol_ma5 if vol_ma5 > 0 else 0
    today_gain = (latest['收盘'] - latest['开盘']) / latest['开盘'] * 100
    c3 = vol_ratio > 1.2 and today_gain > 0
    conditions.append(c3)
    
    return {"触发": sum(conditions) >= 3}


# ============ 最佳策略组合定义 ============

BEST_COMBOS = {
    '深度回调+箱体突破': {
        'strategies': ['深度回调', '箱体突破'],
        'win_rate': 61.5,
        'avg_return': 3.16,
        'description': '回调到支撑位后箱体突破，双重确认'
    },
    '缩量横盘+均线粘合': {
        'strategies': ['缩量横盘', '均线粘合'],
        'win_rate': 57.4,
        'avg_return': 1.19,
        'description': '横盘整理+均线粘合，蓄势待发'
    },
    '缩量横盘+地量见底': {
        'strategies': ['缩量横盘', '地量见底'],
        'win_rate': 53.2,
        'avg_return': 0.92,
        'description': '缩量整理后地量反弹'
    },
    '深度回调+箱体突破+均线粘合': {
        'strategies': ['深度回调', '箱体突破', '均线粘合'],
        'win_rate': 60.0,
        'avg_return': 3.23,
        'description': '三重确认，高胜率精选'
    },
}

STRATEGY_FUNCS = {
    '缩量横盘': strategy_shrink_consolidation,
    '深度回调': strategy_deep_pullback_support,
    '箱体突破': strategy_box_breakout,
    '均线粘合': strategy_ma_convergence_divergence,
    '地量见底': strategy_volume_bottom,
}


def check_combo(df, combo_name):
    """检查是否满足组合条件"""
    combo = BEST_COMBOS[combo_name]
    triggered_strategies = []
    
    for s_name in combo['strategies']:
        func = STRATEGY_FUNCS[s_name]
        result = func(df)
        if result['触发']:
            triggered_strategies.append(s_name)
    
    all_triggered = len(triggered_strategies) == len(combo['strategies'])
    
    return {
        '触发': all_triggered,
        '组合': combo_name,
        '满足策略': triggered_strategies,
        '胜率': combo['win_rate'],
        '平均收益': combo['avg_return'],
        '说明': combo['description']
    }


def get_stock_data(symbol):
    """获取股票数据"""
    try:
        df = ak.stock_zh_a_hist(symbol=symbol, period="daily", 
                                start_date="20241001", adjust="qfq")
        if df is not None and len(df) > 0:
            df['日期'] = pd.to_datetime(df['日期'])
        return df
    except:
        return None


def get_all_stocks():
    """获取全部A股列表"""
    try:
        df = ak.stock_zh_a_spot_em()
        stocks = []
        for _, row in df.iterrows():
            code = row['代码']
            name = row['名称']
            # 过滤ST、退市、北交所
            if 'ST' in name or '退' in name:
                continue
            if code.startswith('8') or code.startswith('4'):
                continue
            stocks.append((code, name, row.get('最新价', 0), row.get('涨跌幅', 0)))
        return stocks
    except Exception as e:
        print(f"获取股票列表失败: {e}")
        return []


def process_stock(args):
    """处理单只股票"""
    code, name, price, change = args
    results = []
    
    try:
        df = get_stock_data(code)
        if df is None or len(df) < 50:
            return results
        
        for combo_name in BEST_COMBOS.keys():
            result = check_combo(df, combo_name)
            if result['触发']:
                results.append({
                    '代码': code,
                    '名称': name,
                    '现价': price,
                    '涨跌幅': change,
                    '组合': combo_name,
                    '历史胜率': result['胜率'],
                    '历史平均收益': result['平均收益'],
                    '说明': result['说明']
                })
        
        if results:
            for r in results:
                safe_print(f"  ★ {code} {name} - {r['组合']} (胜率{r['历史胜率']}%)")
        
        return results
    except:
        return results


def scan_market():
    """扫描全市场"""
    print("=" * 70)
    print("策略组合扫描 - 全市场")
    print("=" * 70)
    print(f"扫描时间: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    
    print("\n最佳策略组合:")
    for name, info in BEST_COMBOS.items():
        print(f"  • {name}: 胜率{info['win_rate']}%, 平均收益{info['avg_return']}%")
    
    print("\n获取股票列表...")
    stocks = get_all_stocks()
    print(f"共{len(stocks)}只股票")
    
    start_time = time.time()
    all_results = []
    completed = 0
    
    print("\n" + "=" * 70)
    print("开始扫描...")
    
    with ThreadPoolExecutor(max_workers=10) as executor:
        futures = {executor.submit(process_stock, stock): stock for stock in stocks}
        
        for future in as_completed(futures):
            completed += 1
            if completed % 200 == 0:
                safe_print(f"\n--- 进度: {completed}/{len(stocks)} ---\n")
            
            try:
                results = future.result()
                all_results.extend(results)
            except:
                pass
    
    elapsed = time.time() - start_time
    print(f"\n扫描完成！耗时: {elapsed:.1f}秒")
    
    return all_results


def print_results(results):
    """打印结果"""
    if not results:
        print("\n今日无符合条件的股票")
        return
    
    print("\n" + "=" * 70)
    print(f"扫描结果: 共{len(results)}个信号")
    print("=" * 70)
    
    # 按组合分类
    by_combo = {}
    for r in results:
        combo = r['组合']
        if combo not in by_combo:
            by_combo[combo] = []
        by_combo[combo].append(r)
    
    # 按胜率排序输出
    sorted_combos = sorted(by_combo.keys(), 
                          key=lambda x: BEST_COMBOS[x]['win_rate'], 
                          reverse=True)
    
    for combo in sorted_combos:
        stocks = by_combo[combo]
        info = BEST_COMBOS[combo]
        
        print(f"\n【{combo}】历史胜率{info['win_rate']}%, 平均收益{info['avg_return']}%")
        print(f"  {info['description']}")
        print("-" * 50)
        
        # 按涨跌幅排序
        stocks_sorted = sorted(stocks, key=lambda x: x['涨跌幅'], reverse=True)
        
        for s in stocks_sorted:
            print(f"  {s['代码']} {s['名称']:<8} 现价{s['现价']:.2f} 涨跌{s['涨跌幅']:+.2f}%")


def export_results(results):
    """导出结果"""
    if not results:
        return
    
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Border, Side
    
    filename = f"策略组合扫描_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    wb = openpyxl.Workbook()
    
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    green_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )
    
    ws = wb.active
    ws.title = "扫描结果"
    
    headers = ['代码', '名称', '现价', '涨跌幅', '策略组合', '历史胜率', '历史平均收益', '说明']
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = border
    
    # 按胜率排序
    results_sorted = sorted(results, 
                           key=lambda x: BEST_COMBOS[x['组合']]['win_rate'], 
                           reverse=True)
    
    for row_idx, r in enumerate(results_sorted, 2):
        ws.cell(row=row_idx, column=1, value=r['代码']).border = border
        ws.cell(row=row_idx, column=2, value=r['名称']).border = border
        ws.cell(row=row_idx, column=3, value=r['现价']).border = border
        
        cell_chg = ws.cell(row=row_idx, column=4, value=f"{r['涨跌幅']:.2f}%")
        cell_chg.border = border
        if r['涨跌幅'] > 0:
            cell_chg.fill = green_fill
        elif r['涨跌幅'] < 0:
            cell_chg.fill = red_fill
        
        ws.cell(row=row_idx, column=5, value=r['组合']).border = border
        ws.cell(row=row_idx, column=6, value=f"{r['历史胜率']}%").border = border
        ws.cell(row=row_idx, column=7, value=f"{r['历史平均收益']}%").border = border
        ws.cell(row=row_idx, column=8, value=r['说明']).border = border
    
    for col in range(1, 9):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 15
    
    wb.save(filename)
    print(f"\n已导出: {filename}")


if __name__ == "__main__":
    results = scan_market()
    print_results(results)
    export_results(results)
