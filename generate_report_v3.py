"""
每日机会分析报告V3 - 改进版
改进点：
1. 加入大盘环境判断，不好就不推荐
2. 排除高位股票
3. 加入"今日不宜操作"的判断
4. 更严格的筛选条件
5. 加入开盘策略（平开下杀的应对）
"""
import requests
import pandas as pd
import numpy as np
import os
from datetime import datetime
import json
import time
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

def log(msg):
    print(msg, flush=True)

for key in list(os.environ.keys()):
    if 'proxy' in key.lower():
        del os.environ[key]

session = requests.Session()
session.trust_env = False
session.headers.update({
    'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36',
})


def get_stock_kline(code, days=60):
    code = str(code).zfill(6)
    kcode = f'sh{code}' if code.startswith('6') else f'sz{code}'
    url = f'https://web.ifzq.gtimg.cn/appstock/app/fqkline/get?param={kcode},day,,,{days},qfq'
    try:
        r = session.get(url, timeout=10)
        data = r.json()
        if not data.get('data'):
            return None
        stock_data = data['data'].get(kcode) or data['data'].get(kcode.upper())
        if not stock_data or 'qfqday' not in stock_data:
            return None
        days_data = stock_data['qfqday']
        df = pd.DataFrame([d[:6] for d in days_data], 
                        columns=['日期','开盘','收盘','最高','最低','成交量'])
        for col in ['开盘','收盘','最高','最低','成交量']:
            df[col] = df[col].astype(float)
        df['日期'] = pd.to_datetime(df['日期'])
        df['涨跌幅'] = df['收盘'].pct_change() * 100
        return df
    except:
        return None


def get_index_kline(code='000001', days=60):
    kcode = f'sh{code}'
    url = f'https://web.ifzq.gtimg.cn/appstock/app/fqkline/get?param={kcode},day,,,{days},qfq'
    try:
        r = session.get(url, timeout=10)
        data = r.json()
        stock_data = data['data'].get(kcode) or data['data'].get(kcode.upper())
        days_data = stock_data.get('qfqday') or stock_data.get('day')
        df = pd.DataFrame([d[:6] for d in days_data], 
                        columns=['日期','开盘','收盘','最高','最低','成交量'])
        for col in ['开盘','收盘','最高','最低','成交量']:
            df[col] = df[col].astype(float)
        df['日期'] = pd.to_datetime(df['日期'])
        df['涨跌幅'] = df['收盘'].pct_change() * 100
        return df
    except:
        return None


def get_stock_name(code):
    code = str(code).zfill(6)
    kcode = f'sh{code}' if code.startswith('6') else f'sz{code}'
    try:
        url = f'https://qt.gtimg.cn/q={kcode}'
        r = session.get(url, timeout=5)
        parts = r.text.split('~')
        if len(parts) > 1:
            return parts[1]
    except:
        pass
    return code


def analyze_market():
    """分析大盘环境"""
    df = get_index_kline('000001', 60)
    if df is None or len(df) < 20:
        return {'score': 0, 'trend': '未知', 'advice': '无法获取大盘数据'}
    
    current = df.iloc[-1]
    prev = df.iloc[-2]
    
    # 均线
    ma5 = df['收盘'].iloc[-5:].mean()
    ma10 = df['收盘'].iloc[-10:].mean()
    ma20 = df['收盘'].iloc[-20:].mean()
    
    # 近期走势
    ret_3d = (current['收盘'] - df.iloc[-4]['收盘']) / df.iloc[-4]['收盘'] * 100
    ret_5d = (current['收盘'] - df.iloc[-6]['收盘']) / df.iloc[-6]['收盘'] * 100
    
    # 量能
    vol_ratio = df['成交量'].iloc[-5:].mean() / df['成交量'].iloc[-20:].mean()
    
    # 评分
    score = 0
    
    # 均线排列
    if current['收盘'] > ma5 > ma10 > ma20:
        score += 40
        trend = '多头排列'
    elif current['收盘'] > ma5 > ma10:
        score += 25
        trend = '偏多'
    elif current['收盘'] < ma5 < ma10:
        score -= 30
        trend = '空头排列'
    else:
        trend = '震荡'
    
    # 近期涨跌
    if ret_5d > 2:
        score += 20
    elif ret_5d < -2:
        score -= 20
    
    # 量能
    if vol_ratio > 1.2:
        score += 10
    elif vol_ratio < 0.8:
        score -= 10
    
    # 昨日走势
    if current['涨跌幅'] > 0.5:
        score += 15
    elif current['涨跌幅'] < -0.5:
        score -= 15
    
    # 生成建议
    if score >= 50:
        advice = '大盘强势，可以积极操作'
        position = '60-80%'
        action = '正常买入'
    elif score >= 20:
        advice = '大盘偏强，可以适度操作'
        position = '40-60%'
        action = '轻仓试探'
    elif score >= 0:
        advice = '大盘震荡，谨慎操作'
        position = '30-50%'
        action = '观望为主'
    elif score >= -20:
        advice = '大盘偏弱，减少操作'
        position = '10-30%'
        action = '只看不动'
    else:
        advice = '大盘弱势，不宜操作'
        position = '空仓'
        action = '休息等待'
    
    return {
        'score': score,
        'trend': trend,
        'advice': advice,
        'position': position,
        'action': action,
        'ret_5d': ret_5d,
        'vol_ratio': vol_ratio,
        'last_change': current['涨跌幅'],
    }


def analyze_stock(code):
    """分析个股"""
    df = get_stock_kline(code, 60)
    if df is None or len(df) < 20:
        return None
    
    name = get_stock_name(code)
    if 'ST' in name or 'st' in name:
        return None
    
    current = df.iloc[-1]
    
    # 均线
    ma5 = df['收盘'].iloc[-5:].mean()
    ma10 = df['收盘'].iloc[-10:].mean()
    ma20 = df['收盘'].iloc[-20:].mean()
    
    # 位置
    high_20d = df['最高'].iloc[-20:].max()
    low_20d = df['最低'].iloc[-20:].min()
    pos_20d = (current['收盘'] - low_20d) / (high_20d - low_20d) * 100 if high_20d > low_20d else 50
    
    # 近期涨幅
    ret_5d = (current['收盘'] - df.iloc[-6]['收盘']) / df.iloc[-6]['收盘'] * 100
    ret_10d = (current['收盘'] - df.iloc[-11]['收盘']) / df.iloc[-11]['收盘'] * 100
    
    # 量能
    vol_ratio = df['成交量'].iloc[-5:].mean() / df['成交量'].iloc[-20:].mean()
    
    # 波动率
    volatility = df['涨跌幅'].iloc[-10:].std()
    
    # 评分
    score = 0
    risks = []
    
    # 均线
    if current['收盘'] > ma5 > ma10 > ma20:
        score += 30
    elif current['收盘'] > ma5 > ma10:
        score += 20
    elif current['收盘'] > ma5:
        score += 10
    elif current['收盘'] < ma5 < ma10:
        score -= 20
        risks.append('均线空头')
    
    # 位置 - 严格排除高位
    if pos_20d > 85:
        score -= 30
        risks.append('高位风险')
    elif pos_20d > 70:
        score -= 10
    elif 30 <= pos_20d <= 70:
        score += 20
    elif pos_20d < 30:
        score += 15
    
    # 近期涨幅 - 不追涨
    if ret_5d > 8:
        score -= 25
        risks.append('涨幅过大')
    elif ret_5d > 5:
        score -= 10
    elif -3 <= ret_5d <= 5:
        score += 15
    elif ret_5d < -5:
        score -= 10
        risks.append('下跌趋势')
    
    # 量能
    if 1.0 <= vol_ratio <= 1.5:
        score += 15
    elif vol_ratio > 2:
        score -= 5
        risks.append('量能过大')
    
    # 波动率
    if volatility > 4:
        risks.append('波动较大')
    
    # 综合评级
    if score >= 50 and len(risks) == 0:
        rating = 'A'
        recommend = '可以买入'
    elif score >= 30 and len(risks) <= 1:
        rating = 'B'
        recommend = '观望为主'
    else:
        rating = 'C'
        recommend = '不建议'
    
    return {
        'code': code,
        'name': name,
        'score': score,
        'rating': rating,
        'recommend': recommend,
        'pos_20d': pos_20d,
        'ret_5d': ret_5d,
        'ret_10d': ret_10d,
        'vol_ratio': vol_ratio,
        'volatility': volatility,
        'risks': risks,
        'price': current['收盘'],
        'last_change': current['涨跌幅'],
    }


def generate_report():
    """生成报告"""
    log("="*70)
    log("生成每日机会分析报告V3")
    log(f"时间: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    log("="*70)
    
    # 1. 分析大盘
    log("\n分析大盘环境...")
    market = analyze_market()
    log(f"  大盘评分: {market['score']}")
    log(f"  趋势: {market['trend']}")
    log(f"  建议: {market['advice']}")
    
    # 2. 扫描个股
    log("\n扫描个股...")
    
    # 主板股票池（非ST）
    stocks = {
        '人工智能': ['002230', '603019', '002410', '000977'],
        '半导体': ['002371', '603986', '002049', '002185'],
        '机器人': ['002747', '002527', '002472', '002008'],
        '新能源车': ['002594', '002074', '000625', '002129'],
        '光伏': ['601012', '002459', '600438', '002506'],
        '消费电子': ['002456', '002241', '002475', '002938'],
        '医药': ['000538', '002821', '600276', '002007'],
        '军工': ['600893', '002013', '600760', '002179'],
        '新材料': ['002080', '603799', '002340', '600516'],
    }
    
    all_stocks = []
    for sector, codes in stocks.items():
        for code in codes:
            result = analyze_stock(code)
            if result:
                result['sector'] = sector
                all_stocks.append(result)
            time.sleep(0.1)
    
    # 排序
    all_stocks.sort(key=lambda x: x['score'], reverse=True)
    
    # A级股票
    a_stocks = [s for s in all_stocks if s['rating'] == 'A']
    b_stocks = [s for s in all_stocks if s['rating'] == 'B']
    
    log(f"\n  A级股票: {len(a_stocks)}只")
    log(f"  B级股票: {len(b_stocks)}只")
    
    # 3. 生成Word报告
    log("\n生成Word报告...")
    
    doc = Document()
    
    # 设置字体
    doc.styles['Normal'].font.name = '微软雅黑'
    doc.styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
    doc.styles['Normal'].font.size = Pt(11)
    
    # 标题
    title = doc.add_heading('每日机会分析报告', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    date_para = doc.add_paragraph()
    date_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    date_run = date_para.add_run(f'分析时间：{datetime.now().strftime("%Y年%m月%d日 %H:%M")}')
    date_run.font.color.rgb = RGBColor(128, 128, 128)
    
    # ==================== 核心结论 ====================
    doc.add_heading('核心结论', level=1)
    
    conclusion = doc.add_paragraph()
    
    # 根据大盘情况给出不同建议
    if market['score'] < 0:
        conclusion.add_run('【今日建议：观望不动】\n\n').bold = True
        conclusion.add_run(f"大盘评分{market['score']}分，{market['trend']}，")
        conclusion.add_run('不建议今日操作。\n')
        conclusion.add_run('即使有好的个股信号，也建议等待大盘企稳后再介入。\n\n')
        conclusion.add_run('仓位建议：').bold = True
        conclusion.add_run(f"{market['position']}\n")
    elif market['score'] < 30:
        conclusion.add_run('【今日建议：谨慎操作】\n\n').bold = True
        conclusion.add_run(f"大盘评分{market['score']}分，{market['trend']}，")
        conclusion.add_run('建议轻仓试探或观望。\n\n')
        conclusion.add_run('仓位建议：').bold = True
        conclusion.add_run(f"{market['position']}\n")
    else:
        conclusion.add_run('【今日建议：可以操作】\n\n').bold = True
        conclusion.add_run(f"大盘评分{market['score']}分，{market['trend']}，")
        conclusion.add_run('可以正常操作。\n\n')
        conclusion.add_run('仓位建议：').bold = True
        conclusion.add_run(f"{market['position']}\n")
    
    # ==================== 大盘分析 ====================
    doc.add_heading('一、大盘环境', level=1)
    
    market_table = doc.add_table(rows=5, cols=2)
    market_table.style = 'Table Grid'
    
    market_data = [
        ('大盘评分', f"{market['score']}分"),
        ('趋势判断', market['trend']),
        ('5日涨幅', f"{market['ret_5d']:+.2f}%"),
        ('昨日涨跌', f"{market['last_change']:+.2f}%"),
        ('操作建议', market['action']),
    ]
    for i, (key, val) in enumerate(market_data):
        market_table.rows[i].cells[0].text = key
        market_table.rows[i].cells[1].text = str(val)
    
    doc.add_paragraph()
    
    # ==================== 开盘应对策略 ====================
    doc.add_heading('二、开盘应对策略', level=1)
    
    strategy = doc.add_paragraph()
    strategy.add_run('【关键：不要追涨，等回调】\n\n').bold = True
    
    strategy.add_run('1. 高开（>0.5%）\n').bold = True
    strategy.add_run('   - 不要追！等开盘后30分钟回调\n')
    strategy.add_run('   - 回调到开盘价附近再考虑介入\n\n')
    
    strategy.add_run('2. 平开（-0.3%~+0.3%）\n').bold = True
    strategy.add_run('   - 观察前30分钟走势\n')
    strategy.add_run('   - 如果下杀后能收回，可以介入\n')
    strategy.add_run('   - 如果持续下跌，放弃\n\n')
    
    strategy.add_run('3. 低开（<-0.3%）\n').bold = True
    strategy.add_run('   - 今日不操作\n')
    strategy.add_run('   - 等待企稳后再说\n\n')
    
    strategy.add_run('4. 平开下杀的应对\n').bold = True
    strategy.add_run('   - 这是最常见的情况，不要慌\n')
    strategy.add_run('   - 如果30分钟内能收回开盘价，说明有支撑\n')
    strategy.add_run('   - 如果持续下跌，说明今日弱势，放弃\n')
    
    # ==================== 推荐股票 ====================
    doc.add_heading('三、今日推荐', level=1)
    
    if market['score'] < 0:
        no_rec = doc.add_paragraph()
        no_rec.add_run('大盘偏弱，今日不推荐任何股票。\n').bold = True
        no_rec.add_run('请耐心等待大盘企稳后再操作。')
    elif len(a_stocks) == 0:
        no_rec = doc.add_paragraph()
        no_rec.add_run('今日无A级推荐股票。\n').bold = True
        no_rec.add_run('如果一定要操作，可以参考下方B级股票，但需更加谨慎。')
    else:
        for i, s in enumerate(a_stocks[:3], 1):
            doc.add_heading(f'推荐{i}：{s["code"]} {s["name"]}', level=2)
            
            stock_info = doc.add_paragraph()
            stock_info.add_run(f'板块：{s["sector"]}\n')
            stock_info.add_run(f'评分：{s["score"]}分（{s["rating"]}级）\n')
            stock_info.add_run(f'当前价：{s["price"]:.2f}元\n')
            stock_info.add_run(f'20日位置：{s["pos_20d"]:.0f}%\n')
            stock_info.add_run(f'5日涨幅：{s["ret_5d"]:+.1f}%\n')
            stock_info.add_run(f'昨日涨跌：{s["last_change"]:+.1f}%\n\n')
            
            stock_info.add_run('操作建议：\n').bold = True
            stock_info.add_run(f'- 买入时机：开盘后观察30分钟，回调不破分时低点时介入\n')
            stock_info.add_run(f'- 不要追涨：如果高开或快速上涨，放弃\n')
            stock_info.add_run(f'- 仓位：10%以内\n')
            stock_info.add_run(f'- 止损：-3%（必须执行）\n')
            stock_info.add_run(f'- 止盈：+5%减半仓，剩余跟踪\n')
    
    # ==================== B级观察 ====================
    if b_stocks:
        doc.add_heading('四、B级观察（仅供参考）', level=1)
        
        b_table = doc.add_table(rows=min(6, len(b_stocks)+1), cols=5)
        b_table.style = 'Table Grid'
        
        hdr = b_table.rows[0].cells
        hdr[0].text = '代码'
        hdr[1].text = '名称'
        hdr[2].text = '板块'
        hdr[3].text = '20日位置'
        hdr[4].text = '风险提示'
        
        for i, s in enumerate(b_stocks[:5], 1):
            row = b_table.rows[i].cells
            row[0].text = s['code']
            row[1].text = s['name']
            row[2].text = s['sector']
            row[3].text = f"{s['pos_20d']:.0f}%"
            row[4].text = ', '.join(s['risks']) if s['risks'] else '-'
    
    # ==================== 风险提示 ====================
    doc.add_heading('五、风险提示', level=1)
    
    risk = doc.add_paragraph()
    risk.add_run('【必须遵守的纪律】\n\n').bold = True
    risk.add_run('1. 严格止损：单票亏损超过3%必须卖出\n')
    risk.add_run('2. 不要追涨：股票已经涨起来了就不要追\n')
    risk.add_run('3. 控制仓位：单票不超过15%，总仓位按建议控制\n')
    risk.add_run('4. 大盘优先：大盘不好时，个股再好也不做\n')
    risk.add_run('5. 宁可错过，不要做错：没把握就不做\n\n')
    
    risk.add_run('【今日不宜操作的情况】\n\n').bold = True
    risk.add_run('- 大盘低开超过0.5%\n')
    risk.add_run('- 个股低开超过1%\n')
    risk.add_run('- 开盘30分钟内持续下跌无反弹\n')
    risk.add_run('- 成交量异常放大（可能有利空）\n')
    
    # 免责声明
    doc.add_paragraph()
    disclaimer = doc.add_paragraph()
    disclaimer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    disc_run = disclaimer.add_run('【免责声明】本报告仅供参考学习，不构成投资建议。股市有风险，投资需谨慎。')
    disc_run.font.size = Pt(9)
    disc_run.font.color.rgb = RGBColor(128, 128, 128)
    
    # 保存
    filename = f'每日机会分析_{datetime.now().strftime("%Y%m%d_%H%M")}.docx'
    doc.save(filename)
    log(f"\n报告已生成: {filename}")
    
    return filename


if __name__ == "__main__":
    generate_report()
