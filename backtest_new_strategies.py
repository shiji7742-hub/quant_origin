"""
回测5种新洗盘策略（简化版）
直接用股票列表测试
"""
import akshare as ak
import pandas as pd
import numpy as np
from datetime import datetime, timedelta
from concurrent.futures import ThreadPoolExecutor, as_completed
from threading import Lock
import time
from new_washout_strategies import (
    strategy_shrink_consolidation,
    strategy_deep_pullback_support,
    strategy_box_breakout,
    strategy_ma_convergence_divergence,
    strategy_volume_bottom
)

print_lock = Lock()

def safe_print(msg):
    with print_lock:
        print(msg)

# 直接用股票代码测试
TEST_STOCKS = [
    ('000001', '平安银行'), ('600036', '招商银行'), ('601318', '中国平安'),
    ('000858', '五粮液'), ('600519', '贵州茅台'), ('000568', '泸州老窖'),
    ('300750', '宁德时代'), ('002594', '比亚迪'), ('601012', '隆基绿能'),
    ('000725', '京东方A'), ('002475', '立讯精密'), ('600745', '闻泰科技'),
    ('300059', '东方财富'), ('000776', '广发证券'), ('601688', '华泰证券'),
    ('002415', '海康威视'), ('000063', '中兴通讯'), ('002230', '科大讯飞'),
    ('600276', '恒瑞医药'), ('000538', '云南白药'), ('002007', '华兰生物'),
    ('002119', '康强电子'), ('603986', '兆易创新'), ('002049', '紫光国微'),
    ('300124', '汇川技术'), ('002896', '中大力德'), ('300285', '国瓷材料'),
    ('601127', '赛力斯'), ('002074', '国轩高科'), ('300014', '亿纬锂能'),
]

def get_stock_data(symbol, start_date, end_date):
    """获取股票历史数据"""
    try:
        df = ak.stock_zh_a_hist(symbol=symbol, period="daily", 
                                start_date=start_date, end_date=end_date, adjust="qfq")
        if df is not None and len(df) > 0:
            df['日期'] = pd.to_datetime(df['日期'])
        return df
    except:
        return None

def find_strategy_signals(df, strategy_func, strategy_name):
    """在历史数据中找策略信号"""
    if df is None or len(df) < 50:
        return []
    
    signals = []
    df = df.reset_index(drop=True)
    
    for i in range(50, len(df) - 20):
        df_slice = df.iloc[:i+1].copy()
        result = strategy_func(df_slice)
        
        if result['触发']:
            signals.append({
                'idx': i,
                'date': df.iloc[i]['日期'],
                'price': df.iloc[i]['收盘'],
                'detail': result['条件']
            })
    
    return signals

def calculate_returns(df, signal_idx):
    """计算信号后的收益"""
    signal_price = df.iloc[signal_idx]['收盘']
    returns = {}
    
    for days in [5, 10, 20]:
        target_idx = signal_idx + days
        if target_idx < len(df):
            future_price = df.iloc[target_idx]['收盘']
            ret = (future_price - signal_price) / signal_price * 100
            returns[f'{days}日'] = ret
        else:
            returns[f'{days}日'] = None
    
    if signal_idx + 1 < len(df):
        future_data = df.iloc[signal_idx+1:min(signal_idx+21, len(df))]
        if len(future_data) > 0:
            max_price = future_data['最高'].max()
            min_price = future_data['最低'].min()
            returns['最大收益'] = (max_price - signal_price) / signal_price * 100
            returns['最大回撤'] = (min_price - signal_price) / signal_price * 100
    
    return returns

def process_stock(args):
    """处理单只股票"""
    code, name, strategies = args
    results = {s_name: [] for s_name in strategies.keys()}
    
    try:
        df = get_stock_data(code, '20250101', '20260115')
        if df is None or len(df) < 70:
            return results
        
        for s_name, s_func in strategies.items():
            signals = find_strategy_signals(df, s_func, s_name)
            
            for sig in signals:
                returns = calculate_returns(df, sig['idx'])
                ret_10 = returns.get('10日', 0) or 0
                
                results[s_name].append({
                    '代码': code,
                    '名称': name,
                    '信号日期': sig['date'].strftime('%Y-%m-%d'),
                    '买入价': round(sig['price'], 2),
                    **returns
                })
                
                result_mark = "✓" if ret_10 > 0 else "✗"
                safe_print(f"  {result_mark} [{s_name}] {code} {name} {sig['date'].strftime('%m-%d')} 10日:{ret_10:+.1f}%")
        
        return results
    except Exception as e:
        return results

def backtest_all_strategies():
    """回测所有策略"""
    print("=" * 70)
    print("回测5种新洗盘策略")
    print("=" * 70)
    
    strategies = {
        '缩量横盘': strategy_shrink_consolidation,
        '深度回调': strategy_deep_pullback_support,
        '箱体突破': strategy_box_breakout,
        '均线粘合': strategy_ma_convergence_divergence,
        '地量见底': strategy_volume_bottom,
    }
    
    print(f"\n测试股票: {len(TEST_STOCKS)}只")
    print(f"测试策略: {', '.join(strategies.keys())}")
    print(f"回测区间: 2025-01-01 ~ 2026-01-15")
    
    start_time = time.time()
    
    # 收集任务
    all_tasks = [(code, name, strategies) for code, name in TEST_STOCKS]
    
    print(f"\n总任务数: {len(all_tasks)}")
    print("=" * 70)
    print("开始回测...")
    
    # 汇总结果
    all_results = {s_name: [] for s_name in strategies.keys()}
    completed = 0
    
    with ThreadPoolExecutor(max_workers=8) as executor:
        futures = {executor.submit(process_stock, task): task for task in all_tasks}
        
        for future in as_completed(futures):
            completed += 1
            if completed % 10 == 0:
                safe_print(f"\n--- 进度: {completed}/{len(all_tasks)} ---\n")
            
            try:
                results = future.result()
                for s_name, signals in results.items():
                    all_results[s_name].extend(signals)
            except:
                pass
    
    elapsed = time.time() - start_time
    print(f"\n回测完成！耗时: {elapsed:.1f}秒")
    
    return all_results

def analyze_results(all_results):
    """分析各策略结果"""
    print("\n" + "=" * 70)
    print("各策略回测结果对比")
    print("=" * 70)
    
    summary = []
    
    for s_name, signals in all_results.items():
        if not signals:
            print(f"\n【{s_name}】无信号")
            continue
        
        df = pd.DataFrame(signals)
        total = len(df)
        
        valid_10 = df[df['10日'].notna()]['10日']
        if len(valid_10) > 0:
            win_rate = (valid_10 > 0).sum() / len(valid_10) * 100
            avg_ret = valid_10.mean()
            median_ret = valid_10.median()
            max_ret = valid_10.max()
            min_ret = valid_10.min()
            
            wins = valid_10[valid_10 > 0]
            losses = valid_10[valid_10 < 0]
            avg_win = wins.mean() if len(wins) > 0 else 0
            avg_loss = abs(losses.mean()) if len(losses) > 0 else 1
            profit_ratio = avg_win / avg_loss if avg_loss > 0 else 0
            
            print(f"\n【{s_name}】")
            print(f"  信号数: {total}")
            print(f"  10日胜率: {win_rate:.1f}%")
            print(f"  10日平均收益: {avg_ret:.2f}%")
            print(f"  10日中位数: {median_ret:.2f}%")
            print(f"  最大盈利: {max_ret:.2f}%, 最大亏损: {min_ret:.2f}%")
            print(f"  盈亏比: {profit_ratio:.2f}")
            
            summary.append({
                '策略': s_name,
                '信号数': total,
                '10日胜率': win_rate,
                '10日平均': avg_ret,
                '盈亏比': profit_ratio
            })
    
    if summary:
        print("\n" + "=" * 70)
        print("策略排名（按10日胜率）")
        print("=" * 70)
        
        summary_df = pd.DataFrame(summary)
        summary_df = summary_df.sort_values('10日胜率', ascending=False)
        
        for i, row in summary_df.iterrows():
            print(f"  {row['策略']}: 胜率{row['10日胜率']:.1f}%, 平均{row['10日平均']:.2f}%, 盈亏比{row['盈亏比']:.2f}")
    
    return all_results, summary

def export_results(all_results, summary):
    """导出结果"""
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Border, Side
    
    filename = f"新策略回测_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    wb = openpyxl.Workbook()
    
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    green_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )
    
    ws1 = wb.active
    ws1.title = "策略汇总"
    
    headers = ['策略', '信号数', '10日胜率', '10日平均收益', '盈亏比']
    for col, h in enumerate(headers, 1):
        cell = ws1.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = border
    
    for row_idx, s in enumerate(summary, 2):
        ws1.cell(row=row_idx, column=1, value=s['策略']).border = border
        ws1.cell(row=row_idx, column=2, value=s['信号数']).border = border
        ws1.cell(row=row_idx, column=3, value=f"{s['10日胜率']:.1f}%").border = border
        ws1.cell(row=row_idx, column=4, value=f"{s['10日平均']:.2f}%").border = border
        ws1.cell(row=row_idx, column=5, value=f"{s['盈亏比']:.2f}").border = border
    
    for s_name, signals in all_results.items():
        if not signals:
            continue
        
        ws = wb.create_sheet(title=s_name[:10])
        headers = ['代码', '名称', '信号日期', '买入价', '5日', '10日', '20日', '最大收益', '最大回撤']
        
        for col, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.border = border
        
        for row_idx, sig in enumerate(signals, 2):
            row_data = [
                sig['代码'], sig['名称'], sig['信号日期'], sig['买入价'],
                f"{sig.get('5日', 0) or 0:.2f}%",
                f"{sig.get('10日', 0) or 0:.2f}%",
                f"{sig.get('20日', 0) or 0:.2f}%",
                f"{sig.get('最大收益', 0) or 0:.2f}%",
                f"{sig.get('最大回撤', 0) or 0:.2f}%"
            ]
            for col, val in enumerate(row_data, 1):
                cell = ws.cell(row=row_idx, column=col, value=val)
                cell.border = border
                
                if col in [5, 6, 7, 8]:
                    try:
                        num = float(str(val).replace('%', ''))
                        if num > 0:
                            cell.fill = green_fill
                        elif num < 0:
                            cell.fill = red_fill
                    except:
                        pass
    
    for ws in wb.worksheets:
        for col in range(1, 10):
            ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 12
    
    wb.save(filename)
    print(f"\n已导出: {filename}")

if __name__ == "__main__":
    all_results = backtest_all_strategies()
    all_results, summary = analyze_results(all_results)
    export_results(all_results, summary)
