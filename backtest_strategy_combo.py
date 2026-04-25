"""
策略组合回测 - 测试多策略组合的胜率
找出最佳策略组合
"""
import akshare as ak
import pandas as pd
import numpy as np
from datetime import datetime
from concurrent.futures import ThreadPoolExecutor, as_completed
from threading import Lock
from itertools import combinations
import time
from new_washout_strategies import (
    strategy_shrink_consolidation,
    strategy_deep_pullback_support,
    strategy_box_breakout,
    strategy_ma_convergence_divergence,
    strategy_volume_bottom
)

print_lock = Lock()

def safe_print(msg):
    with print_lock:
        print(msg)

# 测试股票池
TEST_STOCKS = [
    ('000001', '平安银行'), ('600036', '招商银行'), ('601318', '中国平安'),
    ('000858', '五粮液'), ('600519', '贵州茅台'), ('000568', '泸州老窖'),
    ('300750', '宁德时代'), ('002594', '比亚迪'), ('601012', '隆基绿能'),
    ('000725', '京东方A'), ('002475', '立讯精密'), ('600745', '闻泰科技'),
    ('300059', '东方财富'), ('000776', '广发证券'), ('601688', '华泰证券'),
    ('002415', '海康威视'), ('000063', '中兴通讯'), ('002230', '科大讯飞'),
    ('600276', '恒瑞医药'), ('000538', '云南白药'), ('002007', '华兰生物'),
    ('002119', '康强电子'), ('603986', '兆易创新'), ('002049', '紫光国微'),
    ('300124', '汇川技术'), ('002896', '中大力德'), ('300285', '国瓷材料'),
    ('601127', '赛力斯'), ('002074', '国轩高科'), ('300014', '亿纬锂能'),
    ('600030', '中信证券'), ('000651', '格力电器'), ('600887', '伊利股份'),
    ('002352', '顺丰控股'), ('300760', '迈瑞医疗'), ('688981', '中芯国际'),
]

# 策略定义
STRATEGIES = {
    '缩量横盘': strategy_shrink_consolidation,
    '深度回调': strategy_deep_pullback_support,
    '箱体突破': strategy_box_breakout,
    '均线粘合': strategy_ma_convergence_divergence,
    '地量见底': strategy_volume_bottom,
}

def get_stock_data(symbol, start_date, end_date):
    """获取股票历史数据"""
    try:
        df = ak.stock_zh_a_hist(symbol=symbol, period="daily", 
                                start_date=start_date, end_date=end_date, adjust="qfq")
        if df is not None and len(df) > 0:
            df['日期'] = pd.to_datetime(df['日期'])
        return df
    except:
        return None

def check_strategies_at_point(df_slice, strategy_names):
    """检查某个时间点哪些策略触发"""
    triggered = []
    for name in strategy_names:
        func = STRATEGIES[name]
        result = func(df_slice)
        if result['触发']:
            triggered.append(name)
    return triggered

def find_combo_signals(df, combo):
    """找出同时满足组合中所有策略的信号"""
    if df is None or len(df) < 50:
        return []
    
    signals = []
    df = df.reset_index(drop=True)
    
    for i in range(50, len(df) - 20):
        df_slice = df.iloc[:i+1].copy()
        triggered = check_strategies_at_point(df_slice, list(STRATEGIES.keys()))
        
        # 检查是否同时满足组合中的所有策略
        if all(s in triggered for s in combo):
            signals.append({
                'idx': i,
                'date': df.iloc[i]['日期'],
                'price': df.iloc[i]['收盘'],
                'triggered': triggered
            })
    
    return signals

def calculate_returns(df, signal_idx):
    """计算信号后的收益"""
    signal_price = df.iloc[signal_idx]['收盘']
    returns = {}
    
    for days in [5, 10, 20]:
        target_idx = signal_idx + days
        if target_idx < len(df):
            future_price = df.iloc[target_idx]['收盘']
            ret = (future_price - signal_price) / signal_price * 100
            returns[f'{days}日'] = ret
        else:
            returns[f'{days}日'] = None
    
    if signal_idx + 1 < len(df):
        future_data = df.iloc[signal_idx+1:min(signal_idx+21, len(df))]
        if len(future_data) > 0:
            max_price = future_data['最高'].max()
            min_price = future_data['最低'].min()
            returns['最大收益'] = (max_price - signal_price) / signal_price * 100
            returns['最大回撤'] = (min_price - signal_price) / signal_price * 100
    
    return returns

def process_stock_for_combos(args):
    """处理单只股票，找出所有组合的信号"""
    code, name, all_combos = args
    
    # 每个组合的信号列表
    combo_signals = {combo: [] for combo in all_combos}
    
    try:
        df = get_stock_data(code, '20250101', '20260115')
        if df is None or len(df) < 70:
            return combo_signals
        
        df = df.reset_index(drop=True)
        
        # 遍历每个时间点
        for i in range(50, len(df) - 20):
            df_slice = df.iloc[:i+1].copy()
            triggered = check_strategies_at_point(df_slice, list(STRATEGIES.keys()))
            
            # 检查每个组合
            for combo in all_combos:
                if all(s in triggered for s in combo):
                    returns = calculate_returns(df, i)
                    ret_10 = returns.get('10日', 0) or 0
                    
                    combo_signals[combo].append({
                        '代码': code,
                        '名称': name,
                        '信号日期': df.iloc[i]['日期'].strftime('%Y-%m-%d'),
                        '买入价': round(df.iloc[i]['收盘'], 2),
                        '触发策略': '+'.join(combo),
                        **returns
                    })
        
        # 打印进度
        total_signals = sum(len(v) for v in combo_signals.values())
        if total_signals > 0:
            safe_print(f"  {code} {name}: 找到{total_signals}个组合信号")
        
        return combo_signals
    except Exception as e:
        return combo_signals

def backtest_all_combos():
    """回测所有策略组合"""
    print("=" * 70)
    print("策略组合回测 - 找出最佳组合")
    print("=" * 70)
    
    strategy_names = list(STRATEGIES.keys())
    
    # 生成所有2策略组合和3策略组合
    combos_2 = list(combinations(strategy_names, 2))
    combos_3 = list(combinations(strategy_names, 3))
    all_combos = combos_2 + combos_3
    
    print(f"\n测试股票: {len(TEST_STOCKS)}只")
    print(f"单策略: {len(strategy_names)}个")
    print(f"2策略组合: {len(combos_2)}个")
    print(f"3策略组合: {len(combos_3)}个")
    print(f"回测区间: 2025-01-01 ~ 2026-01-15")
    
    start_time = time.time()
    
    # 汇总结果
    all_results = {combo: [] for combo in all_combos}
    completed = 0
    
    print("\n" + "=" * 70)
    print("开始回测...")
    
    # 多线程处理
    tasks = [(code, name, all_combos) for code, name in TEST_STOCKS]
    
    with ThreadPoolExecutor(max_workers=8) as executor:
        futures = {executor.submit(process_stock_for_combos, task): task for task in tasks}
        
        for future in as_completed(futures):
            completed += 1
            if completed % 10 == 0:
                safe_print(f"\n--- 进度: {completed}/{len(tasks)} ---\n")
            
            try:
                combo_signals = future.result()
                for combo, signals in combo_signals.items():
                    all_results[combo].extend(signals)
            except:
                pass
    
    elapsed = time.time() - start_time
    print(f"\n回测完成！耗时: {elapsed:.1f}秒")
    
    return all_results, all_combos

def analyze_combo_results(all_results, all_combos):
    """分析组合结果"""
    print("\n" + "=" * 70)
    print("策略组合回测结果")
    print("=" * 70)
    
    summary = []
    
    for combo in all_combos:
        signals = all_results[combo]
        combo_name = '+'.join(combo)
        
        if len(signals) < 3:  # 信号太少不统计
            continue
        
        df = pd.DataFrame(signals)
        total = len(df)
        
        valid_10 = df[df['10日'].notna()]['10日']
        if len(valid_10) > 0:
            win_rate = (valid_10 > 0).sum() / len(valid_10) * 100
            avg_ret = valid_10.mean()
            
            wins = valid_10[valid_10 > 0]
            losses = valid_10[valid_10 < 0]
            avg_win = wins.mean() if len(wins) > 0 else 0
            avg_loss = abs(losses.mean()) if len(losses) > 0 else 1
            profit_ratio = avg_win / avg_loss if avg_loss > 0 else 0
            
            summary.append({
                '组合': combo_name,
                '策略数': len(combo),
                '信号数': total,
                '10日胜率': win_rate,
                '10日平均': avg_ret,
                '盈亏比': profit_ratio,
                '综合得分': win_rate * 0.4 + avg_ret * 10 + profit_ratio * 10  # 综合评分
            })
    
    if not summary:
        print("没有足够的组合信号进行分析")
        return [], summary
    
    # 按综合得分排序
    summary_df = pd.DataFrame(summary)
    summary_df = summary_df.sort_values('综合得分', ascending=False)
    
    print("\n【2策略组合排名】")
    print("-" * 70)
    combo_2 = summary_df[summary_df['策略数'] == 2]
    for i, row in combo_2.head(10).iterrows():
        print(f"  {row['组合']}: 信号{row['信号数']}个, 胜率{row['10日胜率']:.1f}%, "
              f"平均{row['10日平均']:.2f}%, 盈亏比{row['盈亏比']:.2f}")
    
    print("\n【3策略组合排名】")
    print("-" * 70)
    combo_3 = summary_df[summary_df['策略数'] == 3]
    for i, row in combo_3.head(10).iterrows():
        print(f"  {row['组合']}: 信号{row['信号数']}个, 胜率{row['10日胜率']:.1f}%, "
              f"平均{row['10日平均']:.2f}%, 盈亏比{row['盈亏比']:.2f}")
    
    # 找出最佳组合
    print("\n" + "=" * 70)
    print("【最佳策略组合 TOP 5】")
    print("=" * 70)
    
    # 筛选信号数>=5的组合
    valid_combos = summary_df[summary_df['信号数'] >= 5].head(5)
    for i, row in valid_combos.iterrows():
        print(f"\n  ★ {row['组合']}")
        print(f"    信号数: {row['信号数']}")
        print(f"    10日胜率: {row['10日胜率']:.1f}%")
        print(f"    10日平均收益: {row['10日平均']:.2f}%")
        print(f"    盈亏比: {row['盈亏比']:.2f}")
    
    return all_results, summary

def export_combo_results(all_results, summary):
    """导出组合回测结果"""
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Border, Side
    
    filename = f"策略组合回测_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    wb = openpyxl.Workbook()
    
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    green_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    gold_fill = PatternFill(start_color="FFD700", end_color="FFD700", fill_type="solid")
    border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )
    
    # Sheet1: 组合汇总
    ws1 = wb.active
    ws1.title = "组合排名"
    
    headers = ['组合', '策略数', '信号数', '10日胜率', '10日平均收益', '盈亏比', '综合得分']
    for col, h in enumerate(headers, 1):
        cell = ws1.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = border
    
    # 按综合得分排序
    sorted_summary = sorted(summary, key=lambda x: x['综合得分'], reverse=True)
    
    for row_idx, s in enumerate(sorted_summary, 2):
        ws1.cell(row=row_idx, column=1, value=s['组合']).border = border
        ws1.cell(row=row_idx, column=2, value=s['策略数']).border = border
        ws1.cell(row=row_idx, column=3, value=s['信号数']).border = border
        
        cell_wr = ws1.cell(row=row_idx, column=4, value=f"{s['10日胜率']:.1f}%")
        cell_wr.border = border
        if s['10日胜率'] >= 60:
            cell_wr.fill = green_fill
        
        cell_avg = ws1.cell(row=row_idx, column=5, value=f"{s['10日平均']:.2f}%")
        cell_avg.border = border
        if s['10日平均'] > 0:
            cell_avg.fill = green_fill
        elif s['10日平均'] < 0:
            cell_avg.fill = red_fill
        
        ws1.cell(row=row_idx, column=6, value=f"{s['盈亏比']:.2f}").border = border
        ws1.cell(row=row_idx, column=7, value=f"{s['综合得分']:.1f}").border = border
        
        # 前3名标金色
        if row_idx <= 4:
            for col in range(1, 8):
                ws1.cell(row=row_idx, column=col).fill = gold_fill
    
    # Sheet2: 所有信号明细
    ws2 = wb.create_sheet(title="信号明细")
    headers = ['组合', '代码', '名称', '信号日期', '买入价', '5日', '10日', '20日', '最大收益', '最大回撤']
    
    for col, h in enumerate(headers, 1):
        cell = ws2.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = border
    
    row_idx = 2
    for combo, signals in all_results.items():
        combo_name = '+'.join(combo)
        for sig in signals:
            row_data = [
                combo_name, sig['代码'], sig['名称'], sig['信号日期'], sig['买入价'],
                f"{sig.get('5日', 0) or 0:.2f}%",
                f"{sig.get('10日', 0) or 0:.2f}%",
                f"{sig.get('20日', 0) or 0:.2f}%",
                f"{sig.get('最大收益', 0) or 0:.2f}%",
                f"{sig.get('最大回撤', 0) or 0:.2f}%"
            ]
            for col, val in enumerate(row_data, 1):
                cell = ws2.cell(row=row_idx, column=col, value=val)
                cell.border = border
                
                if col in [6, 7, 8, 9]:
                    try:
                        num = float(str(val).replace('%', ''))
                        if num > 0:
                            cell.fill = green_fill
                        elif num < 0:
                            cell.fill = red_fill
                    except:
                        pass
            row_idx += 1
    
    # 调整列宽
    for ws in wb.worksheets:
        for col in range(1, 12):
            ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 14
    
    wb.save(filename)
    print(f"\n已导出: {filename}")
    return filename

if __name__ == "__main__":
    all_results, all_combos = backtest_all_combos()
    all_results, summary = analyze_combo_results(all_results, all_combos)
    if summary:
        export_combo_results(all_results, summary)
